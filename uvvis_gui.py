"""
uvvis_gui.py - Grafische Oberfläche der UV-Vis-Auswertung (Windows / macOS / Linux).

Start:      python uvvis_gui.py
Selbsttest: python uvvis_gui.py --selftest AUSGABEORDNER [--selftest-isnet]
"""
from __future__ import annotations

import json
import os
import sys
import traceback
from pathlib import Path

APP_NAME = "UVVisTool"
APP_VERSION = "1.2.0"

if "--selftest" in sys.argv:
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import numpy as np  # noqa: E402
import yaml  # noqa: E402
import matplotlib  # noqa: E402

matplotlib.use("QtAgg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.backends.backend_agg import FigureCanvasAgg  # noqa: E402
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg  # noqa: E402
from matplotlib.figure import Figure as MplFigure  # noqa: E402
from PySide6.QtCore import QSettings, Qt, QThread, QTimer, Signal  # noqa: E402
from PySide6.QtCore import QUrl  # noqa: E402
from PySide6.QtGui import (QAction, QActionGroup, QColor, QDesktopServices, QIcon, QImage,  # noqa: E402
                           QKeySequence, QPixmap)
from PySide6.QtWidgets import (QAbstractItemView, QApplication, QCheckBox, QComboBox,  # noqa: E402
                               QDialog, QDialogButtonBox, QDoubleSpinBox, QFileDialog, QFormLayout, QGroupBox, QHBoxLayout,
                               QHeaderView, QLabel, QLineEdit, QMainWindow, QMessageBox,
                               QPlainTextEdit, QProgressDialog, QPushButton, QScrollArea,
                               QSpinBox, QSplitter, QTableWidget, QTableWidgetItem,
                               QTabWidget, QVBoxLayout, QWidget)

import uvvis_core as core  # noqa: E402
import uvvis_images as IMG  # noqa: E402
from uvvis_i18n import T, get_lang, set_lang  # noqa: E402
import uvvis_tabs  # noqa: E402
import uvvis_project as PRJ  # noqa: E402
import uvvis_report as REPORT  # noqa: E402

PREVIEW_DPI = 90
FIGSIZE = (8.0, 6.0)


def resource_path(rel):
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return base / rel


def rgba_to_pixmap(arr, max_h=110):
    arr = np.ascontiguousarray(arr)
    h, w = arr.shape[:2]
    img = QImage(arr.data, w, h, 4 * w, QImage.Format_RGBA8888).copy()
    return QPixmap.fromImage(img).scaledToHeight(max_h, Qt.SmoothTransformation)


def parse_float(txt):
    txt = (txt or "").strip().replace(",", ".")
    if not txt:
        return None
    v = float(txt)
    if v <= 0:
        raise ValueError
    return v


class Task(QThread):
    """Funktion im Hintergrund ausführen (Freistellung, CDXML), GUI bleibt bedienbar."""
    done = Signal(object)
    failed = Signal(str)

    def __init__(self, fn, parent=None):
        super().__init__(parent)
        self.fn = fn

    def run(self):
        try:
            self.done.emit(self.fn())
        except Exception as e:
            self.failed.emit(f"{type(e).__name__}: {e}")


class GuiDragger(core.Dragger):
    """Wie core.Dragger, meldet aber jede Änderung (Ziehen von Inset/Bild/Label, Mausrad)
    an die GUI, damit Vorschau, Serienwechsel und Export dasselbe Layout verwenden."""

    def __init__(self, F, on_change):
        self.on_change = on_change
        super().__init__(F, None)

    def _report(self):
        # erst nach allen anderen Release-Handlern auslesen (Label-Drag von matplotlib)
        QTimer.singleShot(0, lambda: self.on_change(self.F.current_layout()))

    def release(self, ev):
        super().release(ev)
        self._report()

    def scroll(self, ev):
        if super().scroll(ev):
            self._report()

    def key(self, ev):
        pass


DEFAULT_GLOB = {"path_length": 1.0, "cutoff": 1.0, "min_points": 3, "bands": "",
                "baseline": "series", "bl_a": 1050.0, "bl_b": 1100.0, "ymax": 1.0,
                "r2": "legend", "show_err": False, "table": True,
                "size": "half_a4", "w_cm": 16.0, "h_cm": 11.0,
                "fmt_pdf": True, "fmt_svg": True, "fmt_png": True,
                "x_auto": True, "x_min": 200.0, "x_max": 1100.0,
                "bands_add": True, "shoulders": False, "show_excl": False, "inset": True,
                "mark_sh": False}

UNITS = [("mM", "mM"), ("uM", "µM"), ("mg/mL", "mg/mL"), ("M", "M")]


class ConcDialog(QDialog):
    """Konzentrationen je Probe eintragen (wenn sie nicht im Probennamen stehen)."""

    def __init__(self, parent, samples, overrides, default_series):
        super().__init__(parent)
        self.setWindowTitle(T("conc_title"))
        self.samples = samples
        lay = QVBoxLayout(self)
        intro = QLabel(T("conc_intro"))
        intro.setWordWrap(True)
        lay.addWidget(intro)

        unit0 = None
        for s_ in samples:
            o = overrides.get(s_["key"]) or {}
            unit0 = unit0 or o.get("unit") or (s_["parsed"][2] if s_["parsed"] else None)
        row = QHBoxLayout()
        row.addWidget(QLabel(T("unit")))
        self.cmb_unit = QComboBox()
        self.cmb_unit.addItem(T("unit_choose"), None)          # keine stille Vorbelegung
        for key, label in UNITS:
            self.cmb_unit.addItem(label, key)
        self.cmb_unit.setCurrentIndex(max(0, self.cmb_unit.findData(unit0)) if unit0 else 0)
        row.addWidget(self.cmb_unit)
        row.addStretch(1)
        lay.addLayout(row)

        self.table = QTableWidget(len(samples), 3)
        self.table.setHorizontalHeaderLabels([T("col_sample"), T("col_series"), T("col_conc")])
        self.table.verticalHeader().setVisible(False)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        for i, s_ in enumerate(samples):
            o = overrides.get(s_["key"]) or {}
            it = QTableWidgetItem(s_["key"])
            it.setFlags(it.flags() & ~Qt.ItemIsEditable)
            self.table.setItem(i, 0, it)
            series = o.get("series") or (s_["parsed"][0] if s_["parsed"] else default_series)
            self.table.setItem(i, 1, QTableWidgetItem(series))
            c = o.get("conc") or (s_["parsed"][1] if s_["parsed"] and not o.get("skip") else None)
            self.table.setItem(i, 2, QTableWidgetItem(f"{c:g}" if c else ""))
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.table.resizeColumnsToContents()
        lay.addWidget(self.table)

        row = QHBoxLayout()
        row.addWidget(QLabel(T("dil_start")))
        self.sp_start = QDoubleSpinBox(decimals=4, minimum=0.0, maximum=1e6, value=0.0)
        row.addWidget(self.sp_start)
        row.addWidget(QLabel(T("dil_factor")))
        self.sp_fac = QDoubleSpinBox(decimals=3, minimum=1.0, maximum=1000, value=2.0)
        row.addWidget(self.sp_fac)
        b = QPushButton(T("dil_fill"))
        b.clicked.connect(self.fill_dilution)
        row.addWidget(b)
        row.addStretch(1)
        lay.addLayout(row)
        hint = QLabel(T("dil_hint"))
        hint.setStyleSheet("color: gray;")
        lay.addWidget(hint)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.accepted.connect(self.try_accept)
        bb.rejected.connect(self.reject)
        lay.addWidget(bb)
        self.table.setMinimumHeight(36 + 30 * min(len(samples), 12))
        self.resize(600, 300 + 30 * min(len(samples), 12))

    def fill_dilution(self):
        if self.sp_start.value() <= 0:
            QMessageBox.warning(self, T("conc_title"), T("start_missing"))
            return
        rows = sorted({i.row() for i in self.table.selectedIndexes()}) or list(range(self.table.rowCount()))
        c = self.sp_start.value()
        for r in rows:
            self.table.item(r, 2).setText(f"{c:.6g}")
            c /= self.sp_fac.value()

    def try_accept(self):
        any_conc = any((self.table.item(i, 2).text() or "").strip() for i in range(self.table.rowCount()))
        if any_conc and self.cmb_unit.currentData() is None:
            QMessageBox.warning(self, T("conc_title"), T("unit_missing"))
            return
        self.accept()

    def overrides(self):
        unit = self.cmb_unit.currentData()
        out = {}
        for i, s_ in enumerate(self.samples):
            series = (self.table.item(i, 1).text() or "").strip()
            try:
                c = parse_float(self.table.item(i, 2).text())
            except ValueError:
                c = None
            out[s_["key"]] = {"conc": c, "unit": unit, "series": series} if c else {"skip": True}
        return out


class MainWindow(QMainWindow):
    log_signal = Signal(str)

    def __init__(self):
        super().__init__()
        self.settings = QSettings(APP_NAME, APP_NAME)
        set_lang(self.settings.value("lang", "de"))
        self.csv_path = None
        self.groups, self.units = {}, {}
        self.samples, self.conc, self.csv_paths = [], {}, []          # Rohproben und Konzentrations-Overrides
        self.series = {}               # je Serie: mw, structure, photo, bg, layout, assets
        self.current = None
        self.F = None
        self.dragger = None
        self._tasks = []
        self._loading = False
        self.glob = dict(DEFAULT_GLOB)
        self.extra_cache = {}
        self.ov_state, self.fl_state = {}, {}           # Programm startet immer leer
        self._tab_index = 0
        self.project_dir = None
        self._undo, self._redo = [], []
        self._restoring = False
        self._saved_snap = None

        self.fig = MplFigure(figsize=FIGSIZE, dpi=PREVIEW_DPI)
        self.canvas = FigureCanvasQTAgg(self.fig)
        self.canvas.setFixedSize(int(FIGSIZE[0] * PREVIEW_DPI), int(FIGSIZE[1] * PREVIEW_DPI))
        self.timer = QTimer(self, singleShot=True, interval=400)
        self.timer.timeout.connect(self.render)
        self.autosave_timer = QTimer(self, singleShot=True, interval=800)
        self.autosave_timer.timeout.connect(self.autosave)

        self.log_signal.connect(self._append_log)
        core.LOG = IMG.LOG = lambda m="": self.log_signal.emit(str(m))
        icon = resource_path("assets/icon.png")
        if icon.exists():
            self.setWindowIcon(QIcon(str(icon)))
        self.build_ui()
        self.setAcceptDrops(True)
        QTimer.singleShot(0, self._init_undo)
        geo = self.settings.value("geometry")
        if geo:
            self.restoreGeometry(geo)
        else:
            self.resize(1320, 900)

    # ------------------------------------------------------------------ UI
    def build_ui(self):
        self.update_title()
        mb = self.menuBar()
        mb.clear()
        m = mb.addMenu(T("menu_file"))
        m.addAction(QAction(T("act_new_project"), self, shortcut=QKeySequence.New, triggered=self.new_project))
        m.addAction(QAction(T("act_open_project"), self, shortcut=QKeySequence.Open,
                            triggered=self.open_project_dialog))
        self.recent_menu = m.addMenu(T("act_recent"))
        self.fill_recent_menu()
        m.addAction(QAction(T("act_save_project"), self, shortcut=QKeySequence.Save, triggered=self.save_project))
        m.addAction(QAction(T("act_save_copy"), self, shortcut=QKeySequence("Ctrl+Shift+S"),
                            triggered=self.save_project_copy))
        m.addAction(QAction(T("act_show_folder"), self, triggered=self.show_project_folder))
        m.addSeparator()
        m.addAction(QAction(T("act_open"), self, shortcut=QKeySequence("Ctrl+Shift+O"), triggered=self.choose_csv))
        m.addAction(QAction(T("act_conc"), self, triggered=self.edit_concentrations))
        m.addAction(QAction(T("act_restart"), self, triggered=self.restart_file))
        m.addAction(QAction(T("act_export"), self, shortcut="Ctrl+E", triggered=self.export_current))
        m.addSeparator()
        m.addAction(QAction(T("act_quit"), self, shortcut=QKeySequence.Quit, triggered=self.close))
        me = mb.addMenu(T("menu_edit"))
        self.act_undo = QAction(T("act_undo"), self, triggered=self.undo)
        self.act_undo.setShortcuts([QKeySequence.Undo])
        self.act_redo = QAction(T("act_redo"), self, triggered=self.redo)
        self.act_redo.setShortcuts([QKeySequence.Redo, QKeySequence("Ctrl+Shift+Z")])
        me.addAction(self.act_undo)
        me.addAction(self.act_redo)
        self._update_undo_actions()
        ml = mb.addMenu(T("menu_lang"))
        grp = QActionGroup(self)
        for code, name in (("de", "Deutsch"), ("en", "English")):
            act = QAction(name, self, checkable=True, checked=get_lang() == code)
            act.triggered.connect(lambda _=False, c=code: self.switch_lang(c))
            grp.addAction(act)
            ml.addAction(act)
        mh = mb.addMenu(T("menu_help"))
        mh.addAction(QAction(T("act_about"), self, triggered=lambda: QMessageBox.about(
            self, T("act_about"), T("about", v=APP_VERSION))))

        # ---- linke Spalte
        left = QWidget()
        lv = QVBoxLayout(left)

        g = QGroupBox(T("grp_data"))
        f = QFormLayout(g)
        row = QHBoxLayout()
        self.btn_open = QPushButton(T("act_open"))
        self.btn_open.clicked.connect(self.choose_csv)
        row.addWidget(self.btn_open)
        self.lbl_file = QLabel(self.files_label() if self.csv_path else T("no_file"))
        self.lbl_file.setWordWrap(True)
        row.addWidget(self.lbl_file, 1)
        f.addRow(row)
        self.cmb_series = QComboBox()
        self.cmb_series.addItems(list(self.groups))
        if self.current:
            self.cmb_series.setCurrentText(self.current)
        self.cmb_series.currentTextChanged.connect(self.series_changed)
        row = QHBoxLayout()
        row.addWidget(self.cmb_series, 1)
        b = QPushButton(T("act_conc"))
        b.clicked.connect(self.edit_concentrations)
        row.addWidget(b)
        f.addRow(T("series"), row)
        lv.addWidget(g)

        g = QGroupBox(T("grp_compound"))
        f = QFormLayout(g)
        self.ed_mw = QLineEdit()
        self.ed_mw.setPlaceholderText(T("mw_hint_none"))
        self.ed_mw.editingFinished.connect(self.mw_edited)
        f.addRow(T("molar_mass"), self.ed_mw)
        self.lbl_chem = QLabel("")
        self.lbl_chem.setWordWrap(True)
        f.addRow("", self.lbl_chem)
        row = QHBoxLayout()
        self.lbl_struct = QLabel("–")
        b1 = QPushButton(T("choose"))
        b1.clicked.connect(self.choose_structure)
        b2 = QPushButton(T("remove"))
        b2.clicked.connect(lambda: self.set_structure(None))
        row.addWidget(self.lbl_struct, 1)
        row.addWidget(b1)
        row.addWidget(b2)
        f.addRow(T("structure"), row)
        self.sp_struct_w = QDoubleSpinBox(decimals=1, minimum=3, maximum=90, singleStep=1, suffix=" %")
        self.sp_struct_w.valueChanged.connect(lambda v: self.image_size_changed("struktur", v))
        f.addRow(T("img_size"), self.sp_struct_w)
        row = QHBoxLayout()
        self.lbl_photo = QLabel("–")
        b1 = QPushButton(T("choose"))
        b1.clicked.connect(self.choose_photo)
        b2 = QPushButton(T("remove"))
        b2.clicked.connect(lambda: self.set_photo(None))
        row.addWidget(self.lbl_photo, 1)
        row.addWidget(b1)
        row.addWidget(b2)
        f.addRow(T("photo"), row)
        self.sp_photo_w = QDoubleSpinBox(decimals=1, minimum=2, maximum=60, singleStep=0.5, suffix=" %")
        self.sp_photo_w.valueChanged.connect(lambda v: self.image_size_changed("kuevette", v))
        f.addRow(T("img_size"), self.sp_photo_w)
        self.cmb_bg = QComboBox()
        for key in ("auto", "isnet", "grabcut", "border", "none"):
            self.cmb_bg.addItem(T({"auto": "bg_auto", "isnet": "bg_isnet", "grabcut": "bg_grabcut",
                                   "border": "bg_border", "none": "bg_off"}[key]), key)
        self.cmb_bg.currentIndexChanged.connect(self.bg_changed)
        f.addRow(T("bg_method"), self.cmb_bg)
        self.lbl_thumb = QLabel()
        self.lbl_thumb.setAlignment(Qt.AlignCenter)
        f.addRow("", self.lbl_thumb)
        lv.addWidget(g)

        g = QGroupBox(T("grp_eval"))
        f = QFormLayout(g)
        self.sp_d = QDoubleSpinBox(decimals=3, minimum=0.001, maximum=100, singleStep=0.1)
        f.addRow(T("path_length"), self.sp_d)
        self.sp_cut = QDoubleSpinBox(decimals=2, minimum=0.05, maximum=5, singleStep=0.1)
        f.addRow(T("cutoff"), self.sp_cut)
        self.sp_min = QSpinBox(minimum=2, maximum=20)
        f.addRow(T("min_points"), self.sp_min)
        self.ed_bands = QLineEdit()
        self.ed_bands.setPlaceholderText(T("wavelengths_ph"))
        f.addRow(T("wavelengths"), self.ed_bands)
        self.cb_bands_add = QCheckBox(T("bands_add"))
        self.cb_shoulders = QCheckBox(T("find_shoulders"))
        f.addRow(self.cb_bands_add)
        f.addRow(self.cb_shoulders)
        self.cmb_bl = QComboBox()
        for key in ("series", "simple", "manual", "off"):
            self.cmb_bl.addItem(T({"series": "bl_series", "simple": "bl_simple", "manual": "bl_manual",
                                   "off": "bl_off"}[key]), key)
        f.addRow(T("baseline"), self.cmb_bl)
        row = QHBoxLayout()
        self.sp_bla = QDoubleSpinBox(decimals=0, minimum=100, maximum=3500, suffix=" nm")
        self.sp_blb = QDoubleSpinBox(decimals=0, minimum=100, maximum=3500, suffix=" nm")
        row.addWidget(self.sp_bla)
        row.addWidget(QLabel("–"))
        row.addWidget(self.sp_blb)
        f.addRow("", row)
        lv.addWidget(g)

        g = QGroupBox(T("grp_plot"))
        f = QFormLayout(g)
        self.sp_ymax = QDoubleSpinBox(decimals=2, minimum=0.05, maximum=10, singleStep=0.1)
        f.addRow(T("ymax"), self.sp_ymax)
        row = QHBoxLayout()
        self.cb_xauto = QCheckBox(T("xrange_auto"))
        self.sp_xmin = QDoubleSpinBox(decimals=0, minimum=100, maximum=3500, singleStep=10)
        self.sp_xmax = QDoubleSpinBox(decimals=0, minimum=100, maximum=3500, singleStep=10)
        row.addWidget(self.cb_xauto)
        row.addWidget(self.sp_xmin)
        row.addWidget(QLabel("–"))
        row.addWidget(self.sp_xmax)
        f.addRow(T("xrange"), row)
        self.cmb_r2 = QComboBox()
        for key in ("legend", "label", "off"):
            self.cmb_r2.addItem(T({"legend": "r2_legend", "label": "r2_label", "off": "r2_off"}[key]), key)
        f.addRow(T("r2_mode"), self.cmb_r2)
        self.cb_err = QCheckBox(T("show_err"))
        self.cb_marksh = QCheckBox(T("mark_sh"))
        self.cb_inset = QCheckBox(T("inset_show"))
        self.cb_tab = QCheckBox(T("inset_table"))
        f.addRow(self.cb_err)
        f.addRow(self.cb_marksh)
        f.addRow(self.cb_inset)
        f.addRow(self.cb_tab)
        self.cb_excl = QCheckBox(T("show_excl"))
        f.addRow(self.cb_excl)
        self.cmb_size = QComboBox()
        for key in ("half_a4", "custom", "std"):
            self.cmb_size.addItem(T({"half_a4": "size_half_a4", "custom": "size_custom",
                                     "std": "size_std"}[key]), key)
        f.addRow(T("size"), self.cmb_size)
        row = QHBoxLayout()
        self.sp_w = QDoubleSpinBox(decimals=1, minimum=4, maximum=40, singleStep=0.5, suffix=" cm")
        self.sp_h = QDoubleSpinBox(decimals=1, minimum=3, maximum=40, singleStep=0.5, suffix=" cm")
        row.addWidget(self.sp_w)
        row.addWidget(QLabel("×"))
        row.addWidget(self.sp_h)
        f.addRow(T("size_wh"), row)
        row = QHBoxLayout()
        self.cb_pdf, self.cb_svg, self.cb_png = QCheckBox("PDF"), QCheckBox("SVG"), QCheckBox("PNG")
        for cb in (self.cb_pdf, self.cb_svg, self.cb_png):
            row.addWidget(cb)
        row.addStretch(1)
        f.addRow(T("formats"), row)
        lv.addWidget(g)

        row = QHBoxLayout()
        b = QPushButton(T("btn_reset_layout"))
        b.clicked.connect(self.reset_layout)
        row.addWidget(b)
        self.btn_export = QPushButton(T("btn_export"))
        self.btn_export.clicked.connect(self.export_dialog)
        self.btn_export.setDefault(True)
        row.addWidget(self.btn_export)
        lv.addLayout(row)
        hint = QLabel(T("drag_hint"))
        hint.setWordWrap(True)
        hint.setStyleSheet("color: gray;")
        lv.addWidget(hint)
        lv.addStretch(1)

        lscroll = QScrollArea()
        lscroll.setWidget(left)
        lscroll.setWidgetResizable(True)
        lscroll.setMinimumWidth(380)

        # ---- rechte Seite
        cscroll = QScrollArea()
        holder = QWidget()
        hl = QHBoxLayout(holder)
        hl.addWidget(self.canvas, 0, Qt.AlignCenter)
        cscroll.setWidget(holder)
        cscroll.setWidgetResizable(True)
        self.tabs = QTabWidget()
        self.table = QTableWidget(0, 7)
        self.table.setHorizontalHeaderLabels([T("col_lam"), T("col_n"), T("col_coeff"),
                                              T("col_intercept"), T("col_r2"), T("col_cut"), T("col_warn")])
        self.table.horizontalHeader().setSectionResizeMode(6, QHeaderView.Stretch)
        self.table.verticalHeader().setVisible(False)
        self.table.setEditTriggers(QTableWidget.DoubleClicked | QTableWidget.EditKeyPressed
                                   | QTableWidget.SelectedClicked)
        self.table.setToolTip(T("cut_tip"))
        self.table.itemChanged.connect(self.cutoff_edited)
        self._row_keys = []
        old_log = self.log.toPlainText() if hasattr(self, "log") else ""
        self.log = QPlainTextEdit(readOnly=True)
        self.log.setPlainText(old_log)
        self.tabs.addTab(self.table, T("tab_results"))
        self.tabs.addTab(self.log, T("tab_log"))
        right = QSplitter(Qt.Vertical)
        right.addWidget(cscroll)
        right.addWidget(self.tabs)
        right.setSizes([650, 220])
        split = QSplitter(Qt.Horizontal)
        split.addWidget(lscroll)
        split.addWidget(right)
        split.setSizes([400, 920])
        self.main_tabs = QTabWidget()
        self.main_tabs.addTab(split, T("tab_eps"))
        self.ov_tab = uvvis_tabs.OverlayTab(self, self.ov_state)
        self.fl_tab = uvvis_tabs.FluoTab(self, self.fl_state)
        self.main_tabs.addTab(self.ov_tab, T("tab_overlay"))
        self.main_tabs.addTab(self.fl_tab, T("tab_fluo"))
        self.main_tabs.setCurrentIndex(self._tab_index)
        self.main_tabs.currentChanged.connect(lambda i: setattr(self, "_tab_index", i))
        self.setCentralWidget(self.main_tabs)

        self.apply_globals()
        self.apply_series_widgets()
        for w in (self.sp_d, self.sp_cut, self.sp_ymax, self.sp_bla, self.sp_blb, self.sp_w, self.sp_h,
                  self.sp_xmin, self.sp_xmax):
            w.valueChanged.connect(self.changed)
        self.cb_xauto.toggled.connect(self._xauto_toggled)
        self.cmb_size.currentIndexChanged.connect(self._size_changed)
        for cb in (self.cb_pdf, self.cb_svg, self.cb_png):
            cb.toggled.connect(lambda *_: self.collect_globals())
        self.sp_min.valueChanged.connect(self.changed)
        self.ed_bands.editingFinished.connect(self.changed)
        for w in (self.cmb_bl, self.cmb_r2):
            w.currentIndexChanged.connect(self.changed)
        for w in (self.cb_err, self.cb_tab, self.cb_bands_add, self.cb_shoulders, self.cb_excl, self.cb_inset,
                  self.cb_marksh):
            w.toggled.connect(self.changed)

    def apply_globals(self):
        self._loading = True
        gl = self.glob
        self.sp_d.setValue(gl["path_length"])
        self.sp_cut.setValue(gl["cutoff"])
        self.sp_min.setValue(gl["min_points"])
        self.ed_bands.setText(gl["bands"])
        self.cmb_bl.setCurrentIndex(self.cmb_bl.findData(gl["baseline"]))
        self.sp_bla.setValue(gl["bl_a"])
        self.sp_blb.setValue(gl["bl_b"])
        self.sp_ymax.setValue(gl["ymax"])
        self.cmb_r2.setCurrentIndex(self.cmb_r2.findData(gl["r2"]))
        self.cb_err.setChecked(gl["show_err"])
        self.cb_tab.setChecked(gl["table"])
        self.cmb_size.setCurrentIndex(max(0, self.cmb_size.findData(gl["size"])))
        self.sp_w.setValue(gl["w_cm"])
        self.sp_h.setValue(gl["h_cm"])
        self.cb_pdf.setChecked(gl["fmt_pdf"])
        self.cb_svg.setChecked(gl["fmt_svg"])
        self.cb_png.setChecked(gl["fmt_png"])
        self.cb_xauto.setChecked(gl["x_auto"])
        self.cb_bands_add.setChecked(gl.get("bands_add", True))
        self.cb_shoulders.setChecked(gl.get("shoulders", False))
        self.cb_excl.setChecked(gl.get("show_excl", False))
        self.cb_inset.setChecked(gl.get("inset", True))
        self.cb_marksh.setChecked(gl.get("mark_sh", False))
        self.cb_tab.setEnabled(gl.get("inset", True))
        self.sp_xmin.setValue(gl["x_min"])
        self.sp_xmax.setValue(gl["x_max"])
        self.sp_xmin.setEnabled(not gl["x_auto"])
        self.sp_xmax.setEnabled(not gl["x_auto"])
        self._update_size_widgets()
        manual = gl["baseline"] == "manual"
        self.sp_bla.setEnabled(manual)
        self.sp_blb.setEnabled(manual)
        self._loading = False

    def collect_globals(self):
        self.glob.update(path_length=self.sp_d.value(), cutoff=self.sp_cut.value(),
                         min_points=self.sp_min.value(), bands=self.ed_bands.text().strip(),
                         baseline=self.cmb_bl.currentData(), bl_a=self.sp_bla.value(),
                         bl_b=self.sp_blb.value(), ymax=self.sp_ymax.value(),
                         r2=self.cmb_r2.currentData(), show_err=self.cb_err.isChecked(),
                         table=self.cb_tab.isChecked(), size=self.cmb_size.currentData(),
                         w_cm=self.sp_w.value(), h_cm=self.sp_h.value(),
                         fmt_pdf=self.cb_pdf.isChecked(), fmt_svg=self.cb_svg.isChecked(),
                         fmt_png=self.cb_png.isChecked(), x_auto=self.cb_xauto.isChecked(),
                         bands_add=self.cb_bands_add.isChecked(), shoulders=self.cb_shoulders.isChecked(),
                         show_excl=self.cb_excl.isChecked(), inset=self.cb_inset.isChecked(),
                         mark_sh=self.cb_marksh.isChecked(),
                         x_min=self.sp_xmin.value(), x_max=self.sp_xmax.value())

    def _update_size_widgets(self):
        key = self.cmb_size.currentData()
        if key == "half_a4":
            self.sp_w.blockSignals(True)
            self.sp_h.blockSignals(True)
            self.sp_w.setValue(16.0)
            self.sp_h.setValue(11.0)
            self.sp_w.blockSignals(False)
            self.sp_h.blockSignals(False)
        custom = key == "custom"
        self.sp_w.setEnabled(custom)
        self.sp_h.setEnabled(custom)

    def apply_series_widgets(self):
        self._loading = True
        st = self.series.get(self.current)
        self.ed_mw.setText(f"{st['mw']:g}" if st and st["mw"] else "")
        self.lbl_struct.setText(Path(st["structure"]).name if st and st["structure"] else "–")
        self.lbl_photo.setText(Path(st["photo"]).name if st and st["photo"] else "–")
        self.cmb_bg.setCurrentIndex(self.cmb_bg.findData(st["bg"] if st else "auto"))
        info = st.get("chem_info", "") if st else ""
        unit = self.units.get(self.current) if st else None
        warn = bool(st and st.get("chem_warn"))
        if unit and unit != "mg/mL":
            info = (info + "\n" if info else "") + T("mw_unused", u="µM" if unit == "uM" else unit)
        mw_missing = unit == "mg/mL" and not (st and st["mw"])
        if mw_missing:
            info = (info + "\n" if info else "") + T("mw_needed")
            warn = True
        self.ed_mw.setStyleSheet("border: 2px solid #d9534f;" if mw_missing else "")
        self.lbl_chem.setText(info)
        self.lbl_chem.setStyleSheet("color: #d9534f;" if warn else "")
        pa = st.get("photo_asset") if st else None
        self.lbl_thumb.setPixmap(rgba_to_pixmap(pa["rgba"]) if pa else QPixmap())
        self.sp_struct_w.setValue(st.get("struct_w", 17.0) if st else 17.0)
        self.sp_photo_w.setValue(st.get("photo_w", 9.0) if st else 9.0)
        self.sp_struct_w.setEnabled(bool(st and st["structure"]))
        self.sp_photo_w.setEnabled(bool(st and st["photo"]))
        self._loading = False

    def _xauto_toggled(self, on):
        if self._loading:
            return
        self.sp_xmin.setEnabled(not on)
        self.sp_xmax.setEnabled(not on)
        if not on and self.current in self.groups:       # mit dem Datenbereich vorbelegen
            xs = np.concatenate([s_["x"] for s_ in self.groups[self.current]])
            self._loading = True
            if self.sp_xmin.value() <= xs.min() or self.sp_xmin.value() >= xs.max():
                self.sp_xmin.setValue(float(np.floor(xs.min())))
            if self.sp_xmax.value() >= xs.max() or self.sp_xmax.value() <= self.sp_xmin.value():
                self.sp_xmax.setValue(float(np.ceil(xs.max())))
            self._loading = False
        self.changed()

    def _size_changed(self, *_):
        if self._loading:
            return
        self._update_size_widgets()
        self.changed()

    def switch_lang(self, code):
        self.collect_globals()
        set_lang(code)
        self.settings.setValue("lang", code)
        self.build_ui()
        self.render()

    def _append_log(self, msg):
        self.log.appendPlainText(msg)

    # ------------------------------------------------------------ Daten
    def choose_csv(self):
        if not self.ensure_project():
            return
        start = self.settings.value("last_dir", str(Path.home()))
        ps, _ = QFileDialog.getOpenFileNames(self, T("dlg_open_csv"), start,
                                             f"{T('flt_csv')};;{T('flt_all')}")
        if ps:
            self.open_files(ps)

    def open_csv(self, path, ask=True):
        return self.open_files([path], ask)

    def files_label(self):
        names = [Path(p).name for p in self.csv_paths]
        if len(names) == 1:
            return names[0]
        short = ", ".join(names[:3]) + (", …" if len(names) > 3 else "")
        return T("files_loaded", n=len(names), names=short)

    def open_files(self, paths, ask=True):
        """Eine oder mehrere Messdateien (CSV/TXT/DSW/BSW) gemeinsam öffnen,
        z. B. wenn jede Konzentration einzeln gemessen wurde."""
        if not self.ensure_project():
            return False
        try:
            paths = sorted(self.import_file(p, "data") for p in paths)   # Kopie ins Projekt
        except OSError as e:
            self.error(str(e))
            return False
        path = paths[0]
        try:
            samples = core.read_samples(paths)
        except Exception as e:
            self.error(f"{type(e).__name__}: {e}")
            return False
        if not samples:
            self.error(T("conc_none"))
            return False
        self.csv_path, self.csv_paths, self.samples = path, paths, samples
        self.settings.setValue("last_dir", str(Path(path).parent))
        self.groups, self.units, self.series, self.conc = {}, {}, {}, {}
        _, _, missing = core.group_samples(samples, self.conc, Path(path).stem)
        if missing and ask and not self.edit_concentrations(rebuild=False):
            pass
        if not self.rebuild_groups():
            self.lbl_file.setText(self.files_label())
            self.error(T("conc_none"))
            return False
        self.lbl_file.setText(self.files_label())
        self.apply_globals()
        self.apply_series_widgets()
        self.render()
        return True

    def edit_concentrations(self, rebuild=True):
        if not self.samples:
            self.error(T("load_first"))
            return False
        dlg = ConcDialog(self, self.samples, self.conc, Path(self.csv_path).stem)
        if dlg.exec() != QDialog.Accepted:
            return False
        self.conc = dlg.overrides()
        if rebuild:
            self.rebuild_groups()
            self.apply_series_widgets()
            self.render()
        return True

    def rebuild_groups(self):
        """Serien aus Proben + Konzentrationen neu bilden; Einstellungen je Serie bleiben."""
        try:
            groups, units, missing = core.group_samples(self.samples, self.conc, Path(self.csv_path).stem)
        except ValueError as e:
            self.error(str(e))
            return False
        for k in missing:
            self.log_signal.emit("  " + T("no_conc_in_name", name=k))
        self.groups, self.units = groups, units
        for g in groups:
            self.series.setdefault(g, {"mw": None, "structure": None, "photo": None, "bg": "auto",
                                       "layout": {}, "structure_asset": None, "photo_asset": None,
                                       "chem_info": "", "chem_warn": False,
                                       "struct_w": 17.0, "photo_w": 9.0, "band_cutoff": {},
                                       "band_hidden": []})
        if self.current not in groups:
            self.current = next(iter(groups), None)
        self.cmb_series.blockSignals(True)
        self.cmb_series.clear()
        self.cmb_series.addItems(list(groups))
        if self.current:
            self.cmb_series.setCurrentText(self.current)
        self.cmb_series.blockSignals(False)
        log_names = ", ".join(groups) or "–"
        self.log_signal.emit(T("cary_series", name=Path(self.csv_path).name, s=log_names))
        return bool(groups)

    def restart_file(self):
        """Eingaben zu den aktuellen Messdateien verwerfen (Konzentrationen, Serien-Einstellungen)."""
        if not self.csv_path:
            return self.error(T("load_first"))
        if os.environ.get("QT_QPA_PLATFORM") != "offscreen":
            ans = QMessageBox.question(self, T("act_restart").rstrip("…"),
                                       T("restart_q", name=self.files_label()))
            if ans != QMessageBox.Yes:
                return False
        paths = list(self.csv_paths)
        self.glob = dict(DEFAULT_GLOB)
        self.series, self.conc, self.current = {}, {}, None
        return self.open_files(paths)

    # ------------------------------------------------------------ Projekt
    def get_state(self):
        """Gesamter Zustand aller Tabs (serialisierbar) – für Projektdatei und Undo."""
        if hasattr(self, "sp_d"):
            self.collect_globals()
        keys = ("mw", "structure", "photo", "bg", "layout", "struct_w", "photo_w", "band_cutoff", "band_hidden")
        return {"eps": {"files": list(self.csv_paths), "glob": dict(self.glob), "conc": self.conc,
                        "series": {g: {k: st.get(k) for k in keys} for g, st in self.series.items()},
                        "current": self.current},
                "overlay": self.ov_state, "fluo": self.fl_state, "tab": self._tab_index}

    def set_state(self, state):
        """Zustand übernehmen (Projekt laden, Undo/Redo)."""
        state = json.loads(json.dumps(state))          # entkoppeln
        eps = state.get("eps") or {}
        old_series = self.series
        self.glob = dict(DEFAULT_GLOB, **(eps.get("glob") or {}))
        self.conc = eps.get("conc") or {}
        self.csv_paths = [p for p in eps.get("files") or [] if Path(p).exists()]
        self.csv_path = self.csv_paths[0] if self.csv_paths else None
        self.groups, self.units, self.series, self.samples = {}, {}, {}, []
        self.ov_state.clear()
        self.ov_state.update(state.get("overlay") or {})
        self.fl_state.clear()
        self.fl_state.update(state.get("fluo") or {})
        self._tab_index = state.get("tab", 0) or 0
        self.current = eps.get("current")
        self.build_ui()
        if self.csv_paths:
            key = ("samples_multi", tuple(self.csv_paths))
            if key not in self.extra_cache:
                self.extra_cache[key] = core.read_samples(self.csv_paths)
            self.samples = self.extra_cache[key]
            self.rebuild_groups()
            for g, sd in (eps.get("series") or {}).items():
                if g not in self.series:
                    continue
                st = self.series[g]
                for k, v in sd.items():
                    if v is not None and k not in ("structure", "photo"):
                        st[k] = v
                st["layout"] = st.get("layout") or {}
                old = old_series.get(g) or {}
                for kind, setter in (("structure", self.set_structure), ("photo", self.set_photo)):
                    p = sd.get(kind)
                    if p and Path(p).exists():
                        if old.get(kind) == p and old.get(kind + "_asset") is not None:
                            st[kind], st[kind + "_asset"] = p, old[kind + "_asset"]   # wiederverwenden
                            if kind == "structure":
                                st["chem_info"], st["chem_warn"] = old.get("chem_info", ""), old.get("chem_warn")
                        elif kind == "structure":
                            setter(p, g, render=False, sync=True, autofill_mw=False)
                        else:
                            setter(p, g, render=False, sync=True)
            if self.current not in self.groups:
                self.current = next(iter(self.groups), None)
            self.cmb_series.blockSignals(True)
            self.cmb_series.setCurrentText(self.current or "")
            self.cmb_series.blockSignals(False)
            self.lbl_file.setText(self.files_label())
        self.apply_globals()
        self.apply_series_widgets()
        if self.csv_paths:
            self.render()
        else:
            self.fig.clear()
            self.canvas.draw_idle()

    def _snap(self):
        return json.dumps(self.get_state(), sort_keys=True, default=str)

    def _init_undo(self):
        self._undo, self._redo = [self._snap()], []
        self._saved_snap = self._undo[0]
        self._update_undo_actions()
        self.update_title()

    def push_undo(self):
        if self._restoring or not self._undo:
            return
        snap = self._snap()
        if snap != self._undo[-1]:
            self._undo.append(snap)
            del self._undo[:-200]
            self._redo.clear()
            self._update_undo_actions()
            self.update_title()
            if self.project_dir:
                self.autosave_timer.start()

    def _restore(self, snap):
        self._restoring = True
        self.set_state(json.loads(snap))
        QTimer.singleShot(200, self._end_restore)       # erst nach dem Neuzeichnen aller Tabs

    def _end_restore(self):
        self._restoring = False
        if self._undo:
            self._undo[-1] = self._snap()              # tatsächlichen Zustand übernehmen
        self._update_undo_actions()
        self.update_title()
        if self.project_dir:
            self.autosave_timer.start()                # Undo/Redo-Stand ebenfalls speichern

    def _load_state_into_ui(self, state):
        """Projekt/neues Projekt: Zustand setzen, Undo-Verlauf nach dem Neuzeichnen neu beginnen."""
        self._restoring = True
        self.set_state(state)
        QTimer.singleShot(200, self._after_load)

    def _after_load(self):
        self._restoring = False
        self._init_undo()

    def undo(self):
        if len(self._undo) > 1:
            self._redo.append(self._undo.pop())
            self._restore(self._undo[-1])

    def redo(self):
        if self._redo:
            self._undo.append(self._redo.pop())
            self._restore(self._undo[-1])

    def _update_undo_actions(self):
        if hasattr(self, "act_undo"):
            self.act_undo.setEnabled(len(self._undo) > 1)
            self.act_redo.setEnabled(bool(self._redo))

    def is_dirty(self):
        if self._restoring:
            return False
        return bool(self._undo) and self._undo[-1] != self._saved_snap

    def update_title(self):
        name = Path(self.project_dir).name if self.project_dir else T("no_project_title")
        self.setWindowTitle(f"{name} – {T('app_title')} {APP_VERSION}")

    # -- Projektordner zuerst: Dateien werden beim Laden kopiert, Änderungen automatisch gespeichert
    def ensure_project(self):
        if self.project_dir:
            return True
        if os.environ.get("QT_QPA_PLATFORM") == "offscreen":
            return False
        if QMessageBox.question(self, T("app_title"), T("need_project_q")) != QMessageBox.Yes:
            return False
        return self.choose_project_folder()

    def choose_project_folder(self):
        d = QFileDialog.getExistingDirectory(self, T("choose_project_dir"),
                                             self.settings.value("last_project_dir", str(Path.home())))
        if not d:
            return False
        if PRJ.is_project(d):
            return self.open_project(d, ask=False)
        return self.create_project(d)

    def create_project(self, d):
        try:
            pdir = PRJ.create(d)
        except OSError as e:
            self.error(str(e))
            return False
        self.project_dir = str(pdir)
        self.settings.setValue("last_project_dir", str(pdir.parent))
        self._load_state_into_ui({})
        self.add_recent(pdir)
        self.log_signal.emit(T("project_created", p=pdir))
        return True

    def import_file(self, path, sub):
        if not self.project_dir:
            return str(Path(path).resolve())
        new = PRJ.import_file(self.project_dir, path, sub)
        if Path(new) != Path(path).resolve():
            self.log_signal.emit("  " + T("file_imported", src=Path(path).name, dst=Path(new).relative_to(
                self.project_dir).as_posix()))
        return new

    def maybe_save(self):
        """Vor Neu/Öffnen/Schließen: ausstehendes automatisches Speichern sofort ausführen."""
        if self.autosave_timer.isActive():
            self.autosave_timer.stop()
            self.autosave()
        return True

    def autosave(self):
        if not self.project_dir or self._restoring:
            return
        try:
            PRJ.write(self.project_dir, self.get_state())
            self._saved_snap = self._undo[-1] if self._undo else None
            self.statusBar().showMessage(T("autosaved"), 2000)
        except Exception as e:
            self.log_signal.emit(f"{T('err_title')}: {e}")

    def new_project(self):
        self.maybe_save()
        self.choose_project_folder()

    def open_project_dialog(self):
        self.maybe_save()
        d = QFileDialog.getExistingDirectory(self, T("act_open_project"),
                                             self.settings.value("last_project_dir", str(Path.home())))
        if not d:
            return
        if PRJ.is_project(d):
            self.open_project(d, ask=False)
        elif QMessageBox.question(self, T("app_title"), T("make_project_q", p=d)) == QMessageBox.Yes:
            self.create_project(d)

    def open_project(self, path, ask=True):
        pdir = PRJ.is_project(path)
        if pdir is None:
            self.error(T("no_project", p=path))
            return False
        self.maybe_save()
        try:
            PRJ.backup(pdir)                           # Sicherung: project.uvvis.bak
            state = PRJ.load(pdir)
        except Exception as e:
            self.error(f"{type(e).__name__}: {e}")
            return False
        for m in PRJ.missing_files(state):
            self.log_signal.emit("  " + T("missing_file", p=m))
        self.project_dir = str(pdir)
        self._load_state_into_ui(state)
        self.add_recent(pdir)
        self.settings.setValue("last_project_dir", str(pdir.parent))
        return True

    def save_project(self):
        if not self.project_dir:
            return self.ensure_project()
        self.autosave_timer.stop()
        self.autosave()
        return True

    def save_project_copy(self):
        """Projekt (inkl. aller Dateien) als Kopie in einen neuen Ordner schreiben, z. B. als Zwischenstand."""
        if not self.project_dir:
            return self.error(T("load_first"))
        d = QFileDialog.getExistingDirectory(self, T("act_save_copy"), str(Path(self.project_dir).parent))
        if d:
            return self._write_project(Path(d), switch=False)
        return False

    def _write_project(self, pdir: Path, switch=True):
        try:
            abs_state = PRJ.save(pdir, self.get_state())
        except Exception as e:
            self.error(f"{type(e).__name__}: {e}")
            return False
        if switch:
            self.project_dir = str(pdir)
            self._load_state_into_ui(abs_state)
            self.add_recent(pdir)
        self.log_signal.emit(T("project_saved", p=pdir))
        return True

    def show_project_folder(self):
        if self.project_dir:
            QDesktopServices.openUrl(QUrl.fromLocalFile(self.project_dir))

    def add_recent(self, pdir):
        rec = [str(pdir)] + [r for r in self.recent() if r != str(pdir)]
        self.settings.setValue("recent_projects", json.dumps(rec[:10]))
        self.fill_recent_menu()

    def recent(self):
        try:
            return [r for r in json.loads(self.settings.value("recent_projects", "[]") or "[]")]
        except Exception:
            return []

    def fill_recent_menu(self):
        if not hasattr(self, "recent_menu"):
            return
        self.recent_menu.clear()
        rec = [r for r in self.recent() if PRJ.is_project(r)]
        for r in rec:
            act = QAction(f"{Path(r).name}  –  {Path(r).parent}", self)
            act.triggered.connect(lambda _=False, r=r: self.open_project(r))
            self.recent_menu.addAction(act)
        self.recent_menu.setEnabled(bool(rec))

    def series_changed(self, g):
        if not g or g not in self.series:
            return
        self.current = g
        self.apply_series_widgets()
        self.render()

    def changed(self, *_):
        if self._loading:
            return
        self.collect_globals()
        manual = self.glob["baseline"] == "manual"
        self.sp_bla.setEnabled(manual)
        self.sp_blb.setEnabled(manual)
        self.cb_tab.setEnabled(self.glob.get("inset", True))
        self.timer.start()

    def mw_edited(self):
        if self._loading or not self.current:
            return
        try:
            self.series[self.current]["mw"] = parse_float(self.ed_mw.text())
        except ValueError:
            self.ed_mw.setText("")
            self.series[self.current]["mw"] = None
        self.apply_series_widgets()                 # Warnung zur Molmasse sofort aktualisieren
        self.timer.start()

    # ------------------------------------------------------- Struktur/Foto
    def choose_structure(self):
        if not self.current:
            return self.error(T("load_first"))
        p, _ = QFileDialog.getOpenFileName(self, T("dlg_structure"), self.settings.value("last_img_dir", ""),
                                           f"{T('flt_struct')};;{T('flt_all')}")
        if p:
            self.settings.setValue("last_img_dir", str(Path(p).parent))
            self.set_structure(p)

    def set_structure(self, path, g=None, render=True, sync=False, autofill_mw=True):
        g = g or self.current
        st = self.series[g]
        if path is not None and self.project_dir:
            path = self.import_file(path, "images")
        if path is None:
            st.update(structure=None, structure_asset=None, chem_info="", chem_warn=False)
            self.apply_series_widgets()
            return self.render()

        def work():
            asset = IMG.load_structure(path)
            p = Path(path)
            if asset["mw"] is None and p.suffix.lower() != ".cdxml":
                sib = p.with_suffix(".cdxml")            # gleichnamige CDXML neben SVG/PDF
                if sib.exists():
                    import uvvis_chem
                    chem = uvvis_chem.load_cdxml(sib)
                    asset.update(mw=chem["mw"], formula=chem["formula"], warnings=chem["warnings"])
            return asset

        def done(asset):
            st["structure"], st["structure_asset"] = str(path), asset
            info = []
            if asset.get("formula") and asset.get("mw"):
                info.append(T("mw_formula", f=asset["formula"], mw=asset["mw"]))
                if autofill_mw:
                    st["mw"] = round(asset["mw"], 2)
            if asset.get("drawing_from"):
                info.append(T("drawing_from", name=asset["drawing_from"]))
            info += asset.get("warnings") or []
            st["chem_info"] = "\n".join(info)
            st["chem_warn"] = bool(asset.get("warnings"))
            if g == self.current:
                self.apply_series_widgets()
            if render:
                self.render()

        self.run(work, done, sync)

    def choose_photo(self):
        if not self.current:
            return self.error(T("load_first"))
        p, _ = QFileDialog.getOpenFileName(self, T("dlg_photo"), self.settings.value("last_img_dir", ""),
                                           f"{T('flt_photo')};;{T('flt_all')}")
        if p:
            self.settings.setValue("last_img_dir", str(Path(p).parent))
            self.set_photo(p)

    def bg_changed(self):
        if self._loading or not self.current:
            return
        st = self.series[self.current]
        st["bg"] = self.cmb_bg.currentData()
        if st["photo"]:
            self.set_photo(st["photo"])

    def ensure_model(self):
        """Fragt nach dem Modell-Download. True = ISNet nutzbar."""
        if IMG.model_path(existing_only=True):
            return True
        try:
            import onnxruntime  # noqa: F401
        except ImportError:
            return False
        if os.environ.get("QT_QPA_PLATFORM") == "offscreen" and not getattr(self, "_allow_download", False):
            return False
        if not getattr(self, "_allow_download", False):
            ans = QMessageBox.question(self, T("model_q_title"), T("model_q", mb=IMG.MODEL_SIZE_MB))
            if ans != QMessageBox.Yes:
                return False
        dlg = QProgressDialog(T("downloading"), T("cancel"), 0, 100, self)
        dlg.setWindowModality(Qt.WindowModal)
        dlg.show()

        def prog(done, total):
            dlg.setValue(int(100 * done / max(total, 1)))
            QApplication.processEvents()
            return not dlg.wasCanceled()
        try:
            IMG.download_model(prog)
            return True
        except Exception as e:
            self.log_signal.emit(str(e))
            return False
        finally:
            dlg.close()

    def set_photo(self, path, g=None, render=True, sync=False):
        g = g or self.current
        st = self.series[g]
        if path is not None and self.project_dir:
            path = self.import_file(path, "images")
        if path is None:
            st.update(photo=None, photo_asset=None)
            self.apply_series_widgets()
            return self.render()
        method = st["bg"]
        if method in ("auto", "isnet") and not self.ensure_model():
            method = "grabcut" if method == "isnet" else "auto"
        cache = Path(self.csv_path).parent / ".uvvis_cache" if self.csv_path else None

        def done(asset):
            st["photo"], st["photo_asset"] = str(path), asset
            self.log_signal.emit(f"  {Path(path).name}: {asset.get('bg_method')}")
            if g == self.current:
                self.apply_series_widgets()
            if render:
                self.render()

        self.run(lambda: IMG.load_photo(path, method, cache), done, sync)

    def run(self, fn, done, sync=False):
        if sync:
            try:
                done(fn())
            except Exception as e:
                self.error(str(e))
            return
        QApplication.setOverrideCursor(Qt.WaitCursor)
        self.statusBar().showMessage(T("working"))
        t = Task(fn, self)
        self._tasks.append(t)

        def fin():
            QApplication.restoreOverrideCursor()
            self.statusBar().clearMessage()
            self._tasks.remove(t)
        t.done.connect(done)
        t.failed.connect(self.error)
        t.finished.connect(fin)
        t.start()

    # ------------------------------------------------------------ Plot
    def make_cfg(self, g):
        gl = self.glob
        st = self.series[g]
        cfg = core.deep_merge(core.DEFAULTS, {})
        cfg["_base_dir"] = Path(self.csv_path).parent
        cfg["conc_unit"] = self.units[g]
        cfg["molar_mass_g_mol"] = st["mw"]
        cfg["path_length_cm"] = gl["path_length"]
        cfg["max_abs_fit"] = gl["cutoff"]
        cfg["min_fit_points"] = gl["min_points"]
        try:
            bands = [float(x) for x in gl["bands"].replace(";", ",").split(",") if x.strip()]
        except ValueError:
            bands = []
        cfg["wavelengths_nm"] = bands or "auto"
        cfg["bands_add_auto"] = gl.get("bands_add", True)
        cfg["find_shoulders"] = gl.get("shoulders", False)
        cfg["inset"]["show_excluded"] = gl.get("show_excl", False)
        cfg["inset"]["show"] = gl.get("inset", True)
        cfg["labels"]["mark_shoulders"] = gl.get("mark_sh", False)
        cfg["band_cutoff"] = {int(k): float(v) for k, v in (st.get("band_cutoff") or {}).items()}
        cfg["band_hidden"] = [int(k) for k in (st.get("band_hidden") or [])]
        if gl["baseline"] == "off":
            cfg["baseline_nm"] = None
        elif gl["baseline"] == "manual":
            cfg["baseline_nm"] = [gl["bl_a"], gl["bl_b"]]
        else:
            cfg["baseline_nm"] = "auto"
            cfg["baseline_mode"] = gl["baseline"]
        cfg["ylim"] = [0, gl["ymax"]]
        if not gl["x_auto"] and gl["x_max"] > gl["x_min"]:
            cfg["xlim"] = [gl["x_min"], gl["x_max"]]
        cfg["r2_mode"] = gl["r2"]
        cfg["labels"]["show_error"] = gl["show_err"]
        cfg["inset"]["table"] = gl["table"]
        cfg["figsize_in"] = list(FIGSIZE)
        if gl["size"] != "std":
            core.apply_export_size(cfg, gl["w_cm"], gl["h_cm"])
        fmts = [e for e, on in (("pdf", gl["fmt_pdf"]), ("svg", gl["fmt_svg"]), ("png", gl["fmt_png"])) if on]
        cfg["extra_formats"] = fmts[1:]
        cfg["_main_fmt"] = fmts[0] if fmts else "pdf"
        imgs = []
        if st["structure_asset"]:
            imgs.append({"file": st["structure"], "id": "struktur", "width": st.get("struct_w", 17) / 100,
                         "prefer": "top",
                         "_asset": st["structure_asset"]})
        if st["photo_asset"]:
            imgs.append({"file": st["photo"], "id": "kuevette", "width": st.get("photo_w", 9) / 100,
                         "prefer": "right",
                         "_asset": st["photo_asset"]})
        cfg["images"] = imgs
        return cfg

    def render(self):
        if not self.current or self.current not in self.groups:
            return
        g = self.current
        try:
            cfg = self.make_cfg(g)
            core.setup_fonts(plt, cfg)
            cfg2, spectra, results = core.prepare_series(cfg, self.groups[g], g)
            if self.dragger:
                self.dragger.disconnect()
                self.dragger = None
            wpx, hpx = int(cfg2["figsize_in"][0] * PREVIEW_DPI), int(cfg2["figsize_in"][1] * PREVIEW_DPI)
            if (self.canvas.width(), self.canvas.height()) != (wpx, hpx):
                self.canvas.setFixedSize(wpx, hpx)
                QApplication.processEvents()
            self.F = core.build_figure(plt, cfg2, spectra, results, self.get_layout(g), fig=self.fig)
            self.dragger = GuiDragger(self.F, self.layout_changed)
            self.sync_image_sizes(self.F.current_layout())
            self.canvas.draw_idle()
            self.fill_table(results, cfg2)
            self._last = (cfg2, results)
            self.push_undo()
        except Exception as e:
            self.log_signal.emit(traceback.format_exc())
            self.statusBar().showMessage(f"{T('err_title')}: {e}", 8000)

    def size_key(self):
        gl = self.glob
        return "std" if gl["size"] == "std" else f"{gl['w_cm']:g}x{gl['h_cm']:g}"

    def get_layout(self, g):
        lay = self.series[g]["layout"] or {}
        if lay and not all(isinstance(v, dict) for v in lay.values()):
            lay = {"std": lay}                       # altes Format (eine Größe)
            self.series[g]["layout"] = lay
        return lay.setdefault(self.size_key(), {})

    def layout_changed(self, layout):
        if not self.current or self.F is None:
            return
        stored = self.get_layout(self.current)
        if layout != stored:
            self.series[self.current]["layout"][self.size_key()] = layout
        self.sync_image_sizes(layout)
        self.push_undo()

    def sync_image_sizes(self, layout):
        """Größenfelder an die tatsächliche Bildgröße anpassen (Mausrad, automatisches Verkleinern)."""
        st = self.series.get(self.current)
        if not st:
            return
        self._loading = True
        for key, field, spin in (("image_struktur", "struct_w", self.sp_struct_w),
                                 ("image_kuevette", "photo_w", self.sp_photo_w)):
            if key in layout and len(layout[key]) == 4:
                st[field] = round(layout[key][2] * 100, 1)
                spin.setValue(st[field])
        self._loading = False

    def image_size_changed(self, iid, value):
        """Größe aus dem Zahlenfeld: Breite setzen, Mittelpunkt und Seitenverhältnis behalten."""
        if self._loading or not self.current:
            return
        st = self.series[self.current]
        st["struct_w" if iid == "struktur" else "photo_w"] = value
        lay = self.get_layout(self.current)
        key = f"image_{iid}"
        if key in lay and len(lay[key]) == 4:
            x, y, w, h = lay[key]
            nw = value / 100
            nh = h * nw / w
            nx = min(max(x + (w - nw) / 2, 0.0), max(0.0, 1 - nw))     # in der Achse halten
            ny = min(max(y + (h - nh) / 2, 0.0), max(0.0, 1 - nh))
            lay[key] = [nx, ny, nw, nh]
        self.timer.start()

    def reset_layout(self):
        if self.current:
            self.get_layout(self.current)
            self.series[self.current]["layout"][self.size_key()] = {}
            self.render()

    def fill_table(self, results, cfg):
        sym = "ε [M⁻¹ cm⁻¹]" if core.conc_factor(cfg) else "a [L g⁻¹ cm⁻¹]"
        self.table.blockSignals(True)
        self.table.setHorizontalHeaderLabels([T("col_lam"), T("col_n"), sym, T("col_intercept"),
                                              T("col_r2"), T("col_cut"), T("col_warn")])
        self.table.setRowCount(len(results))
        self._row_keys = []
        for i, r in enumerate(results):
            ft = r["fit"]
            lam = f"{r['lam']:g}" + (" (sh)" if r.get("shoulder") else "")
            cut = r.get("cutoff", cfg["max_abs_fit"])
            if ft:
                v, e = core.fmt_ve(r["eps"], r["eps_err"], max_decimals=core.eps_max_decimals(cfg))
                b, be = core.fmt_ve(ft["intercept"], ft["se_intercept"])
                cells = [lam, str(ft["n"]), f"{v} ± {e}" if e else v,
                         f"{b} ± {be}" if be else b, f"{ft['r2']:.5f}", f"{cut:g}", "; ".join(r["warnings"])]
            else:
                cells = [lam, str(len(r["used"])), "–", "–", "–", f"{cut:g}",
                         T("no_fit", n=len(r["used"]), cut=cut, min=cfg["min_fit_points"])]
            for j, c in enumerate(cells):
                it = QTableWidgetItem(c)
                if j < 6:
                    it.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                if j == 0:                                   # Häkchen = im Bild anzeigen
                    it.setFlags((it.flags() | Qt.ItemIsUserCheckable) & ~Qt.ItemIsEditable)
                    hidden = int(round(r["lam_requested"])) in set(cfg.get("band_hidden") or [])
                    it.setCheckState(Qt.Unchecked if hidden else Qt.Checked)
                elif j != 5:
                    it.setFlags(it.flags() & ~Qt.ItemIsEditable)
                elif abs(cut - cfg["max_abs_fit"]) > 1e-9:
                    it.setBackground(QColor("#fff2cc"))
                    it.setForeground(QColor("#000000"))
                self.table.setItem(i, j, it)
            self._row_keys.append(int(round(r["lam_requested"])))
        self.table.resizeColumnsToContents()
        self.table.blockSignals(False)

    def cutoff_edited(self, item):
        """Cutoff pro Bande aus der Ergebnistabelle (leer = globaler Cutoff); Häkchen in Spalte λ
        = Bande im Bild anzeigen."""
        if not self.current or item.row() >= len(self._row_keys):
            return
        if item.column() == 0:
            st = self.series[self.current]
            hid = set(st.get("band_hidden") or [])
            k = self._row_keys[item.row()]
            hid.discard(k) if item.checkState() == Qt.Checked else hid.add(k)
            st["band_hidden"] = sorted(hid)
            QTimer.singleShot(0, self.render)
            return
        if item.column() != 5:
            return
        key = str(self._row_keys[item.row()])
        st = self.series[self.current]
        bc = st.setdefault("band_cutoff", {})
        try:
            v = parse_float(item.text())
        except ValueError:
            v = None
        if v is None or abs(v - self.glob["cutoff"]) < 1e-9:
            bc.pop(key, None)
        else:
            bc[key] = v
        QTimer.singleShot(0, self.render)

    # ------------------------------------------------------------ Export
    def export_dialog(self):
        if not self.csv_path:
            return self.error(T("load_first"))
        if not self.ensure_project():
            return
        out = self.export_all()
        folders = sorted({str(p.parent) for p in out})
        self.info(T("exported_to", p="\n".join(folders)))

    def export_all(self, folder: Path = None):
        """Alle Serien exportieren: Abbildungen, Ergebnis-CSV und Excel-Report.
        Ohne Ordnerangabe nach <Projekt>/exports/epsilon/<Serie>/ (wird überschrieben)."""
        self.collect_globals()
        stem = Path(self.csv_path).stem
        written = []
        for g in self.groups:
            if folder is None:
                gdir = PRJ.export_dir(self.project_dir, "epsilon", core.safe_name(g))
                name = core.safe_name(g)
            else:
                gdir = Path(folder)
                name = f"{stem}_{g}" if len(self.groups) > 1 else stem
            cfg = self.make_cfg(g)
            core.setup_fonts(plt, cfg)
            cfg2, spectra, results = core.prepare_series(cfg, self.groups[g], g)
            folder_g = gdir
            cfg2["output"] = str(folder_g / f"{name}.{cfg2['_main_fmt']}")
            fig = MplFigure(figsize=cfg2["figsize_in"], dpi=PREVIEW_DPI)
            FigureCanvasAgg(fig)
            F = core.build_figure(plt, cfg2, spectra, results, self.get_layout(g), fig=fig)
            F.export()
            core.write_results(results, folder_g / f"{name}_results.csv", cfg2)
            written += [folder_g / f"{name}.{e}" for e in [cfg2["_main_fmt"]] + cfg2["extra_formats"]]
            written.append(folder_g / f"{name}_results.csv")
            try:
                rep = REPORT.write_report(folder_g / f"{name}_report.xlsx", cfg2, self.groups[g], spectra,
                                          results, g, files=self.csv_paths, version=APP_VERSION)
                written.append(Path(rep))
            except PermissionError:
                self.error(T("report_locked", p=folder_g / f"{name}_report.xlsx"))
        return written

    # ------------------------------------------------------------ Drag & Drop
    def dragEnterEvent(self, ev):
        if ev.mimeData().hasUrls():
            ev.acceptProposedAction()

    def dropEvent(self, ev):
        files = [u.toLocalFile() for u in ev.mimeData().urls()]
        for p in files:
            if PRJ.is_project(p):
                self.open_project(p)
                return
        data = [p for p in files if Path(p).suffix.lower() in {".csv", ".txt"} | core.BINARY_EXT]
        if data:
            self.open_files(data)                     # mehrere Messdateien = ein Datensatz
        for p in files:
            ext = Path(p).suffix.lower()
            if p in data:
                continue
            if not self.current:
                self.error(T("load_first"))
            elif ext in IMG.CHEM_EXT | IMG.VECTOR_EXT:
                self.set_structure(p)
            elif ext in IMG.RASTER_EXT:
                self.set_photo(p)

    # ------------------------------------------------------------ Diverses
    def error(self, msg):
        self.log_signal.emit(f"{T('err_title')}: {msg}")
        if os.environ.get("QT_QPA_PLATFORM") != "offscreen":
            QMessageBox.critical(self, T("err_title"), str(msg))

    def rebuild_tabs(self):
        self.build_ui()
        self.render()

    def export_current(self):
        i = self.main_tabs.currentIndex()
        if i == 1:
            return self.ov_tab.export_dialog()
        if i == 2:
            return self.fl_tab.export_dialog()
        return self.export_dialog()

    def info(self, msg):
        self.log_signal.emit(msg)
        if os.environ.get("QT_QPA_PLATFORM") != "offscreen":
            QMessageBox.information(self, T("app_title"), msg)

    def closeEvent(self, ev):
        if not self.maybe_save():
            ev.ignore()
            return
        self.settings.setValue("geometry", self.saveGeometry())
        super().closeEvent(ev)


# ----------------------------------------------------------------------------
# Selbsttest (für CI und nach der Installation)
# ----------------------------------------------------------------------------
def make_selftest_inputs(d: Path):
    """Synthetische Cary-CSV (2 Serien), CDXML und Küvettenfoto (HEIC, falls möglich)."""
    from PIL import Image, ImageDraw
    d.mkdir(parents=True, exist_ok=True)
    wl = np.arange(1100.0, 199.0, -1.0)
    bands = {"TEST-A": [(262, 6.8, 20), (484, 1.6, 30), (732, 0.25, 40), (816, 0.24, 40)],
             "TEST-B": [(255, 7.3, 18), (351, 2.7, 20), (444, 3.5, 12), (719, 0.42, 35), (804, 0.41, 35)]}
    concs = [("1", 1.0), ("0p5", 0.5), ("0p25", 0.25), ("0p125", 0.125), ("0p0625", 0.0625)]
    rng = np.random.default_rng(0)
    cols, names = [], []
    for ser, bl in bands.items():
        a = sum(h * np.exp(-0.5 * ((wl - l) / w) ** 2) for l, h, w in bl)
        for tag, c in concs:
            names.append(f"{ser}-{tag}mgml_THF")
            y = np.minimum(a * c, 4.0) - 0.03 + rng.normal(0, 0.002, wl.size)
            cols.append(y)
    lines = ["".join(f"{n},," for n in names), "".join("Wavelength (nm),Abs," for _ in names)]
    for i, x in enumerate(wl):
        lines.append("".join(f"{x:.6f},{y[i]:.8f}," for y in cols))
    lines += ["", "TEST-A-1mgml_THF", "Instrument  Cary 60"]
    (d / "selftest.csv").write_text("\r\n".join(lines), encoding="utf-8", newline="")
    nodes = [(1, 100, 100, 6), (2, 112.5, 92.8, 6), (3, 125, 100, 6), (4, 125, 114.4, 6),
             (5, 112.5, 121.6, 6), (6, 100, 114.4, 7), (7, 112.5, 78.4, 8)]
    bonds = [(1, 2, 2), (2, 3, 1), (3, 4, 2), (4, 5, 1), (5, 6, 2), (6, 1, 1), (2, 7, 1)]
    xml = ['<?xml version="1.0" encoding="UTF-8" ?><CDXML BondLength="14.40"><page id="1"><fragment id="99">']
    xml += [f'<n id="{i}" p="{x} {y}"' + (f' Element="{e}"' if e != 6 else "") + "/>" for i, x, y, e in nodes]
    xml += [f'<b id="{20 + k}" B="{b}" E="{e}"' + (' Order="2"' if o == 2 else "") + "/>"
            for k, (b, e, o) in enumerate(bonds)]
    xml += ["</fragment></page></CDXML>"]
    (d / "pyridinol.cdxml").write_text("".join(xml), encoding="utf-8")
    im = Image.new("RGB", (600, 900), (190, 185, 175))
    dr = ImageDraw.Draw(im)
    dr.rectangle([240, 100, 360, 200], fill=(25, 25, 25))
    dr.rectangle([220, 200, 380, 820], fill=(215, 220, 228))
    dr.rectangle([235, 420, 365, 805], fill=(195, 90, 35))
    photo = d / "kuevette.png"
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
        im.save(d / "kuevette.heic", quality=90)
        photo = d / "kuevette.heic"
    except Exception:
        im.save(photo)
    return d / "selftest.csv", d / "pyridinol.cdxml", photo


def selftest(outdir, with_isnet=False):
    out = Path(outdir).resolve()
    out.mkdir(parents=True, exist_ok=True)
    report = [f"{APP_NAME} {APP_VERSION} selftest", f"python {sys.version.split()[0]}, {sys.platform}"]
    ok = True
    app = QApplication.instance() or QApplication(sys.argv[:1])  # noqa: F841
    try:
        csv, cdxml, photo = make_selftest_inputs(out / "input")
        report.append(f"photo format: {photo.suffix}")
        w = MainWindow()
        w._allow_download = with_isnet
        proj = out / "projekt_test"
        assert not w.open_csv(csv), "ohne Projekt darf nichts geladen werden"
        assert w.create_project(proj)
        QApplication.processEvents()
        w.open_csv(csv)
        assert set(w.groups) == {"TEST-A", "TEST-B"}, w.groups
        assert Path(w.csv_path).parent == proj / "data", w.csv_path          # ins Projekt kopiert
        w.set_structure(str(cdxml), "TEST-A", render=False, sync=True)
        st = w.series["TEST-A"]
        report.append(f"CDXML: {st['chem_info']!r}")
        assert st["mw"] and abs(st["mw"] - 95.10) < 0.05, st["mw"]      # 3-Hydroxypyridin C5H5NO
        w.series["TEST-A"]["bg"] = "isnet" if with_isnet else "grabcut"
        w.set_photo(str(photo), "TEST-A", render=False, sync=True)
        bgm = st["photo_asset"]["bg_method"]
        report.append(f"background: {bgm}")
        if with_isnet:
            assert bgm == "isnet", bgm
        w.render()
        files = w.export_all(out)
        for f in files:
            assert f.exists() and f.stat().st_size > 200, f
        import pymupdf
        page = pymupdf.open(str(out / "selftest_TEST-A.pdf"))[0]
        report.append(f"pdf: {len(page.get_drawings())} vector paths, {len(page.get_images())} images")
        assert page.get_images(), "photo missing in PDF"
        svg = (out / "selftest_TEST-A.svg").read_text(encoding="utf-8")
        assert "structure_0" in svg and "<image" in svg, "SVG ohne Struktur/Foto"
        report.append(f"pdf size: {page.rect.width / 72 * 2.54:.2f} x {page.rect.height / 72 * 2.54:.2f} cm")
        # Probennamen ohne Konzentration (z. B. c0..c4) -> Overrides wie aus dem Dialog
        txt = csv.read_text(encoding="utf-8").splitlines()
        n_samples = len([n for n in txt[0].split(",") if n])
        txt[0] = "".join(f"c{i},," for i in range(n_samples))
        plain = out / "input" / "unnamed.csv"
        plain.write_text("\r\n".join(txt), encoding="utf-8", newline="")
        assert w.open_csv(plain, ask=False) is False                 # ohne Angaben: nichts
        w.conc = {f"c{i}": {"conc": 1.0 / 2 ** i, "unit": "mg/mL", "series": "unbenannt"}
                  for i in range(5)}
        assert w.rebuild_groups() and "unbenannt" in w.groups, w.groups
        assert len(w.groups["unbenannt"]) == 5
        w.render()
        report.append("unnamed samples: OK")
        # Einzelmessungen in getrennten Dateien + Cary-Binärdatei (.DSW)
        import struct
        singles = []
        for i, s_ in enumerate(core.read_samples([csv])[:5]):
            f = out / "input" / f"single_{i}.csv"
            f.write_text("\r\n".join([f"{s_['name']},,", "Wavelength (nm),Abs,"] +
                                      [f"{a:.4f},{b:.6f}," for a, b in zip(s_["x"], s_["y"])]),
                         encoding="utf-8", newline="")
            singles.append(f)
        assert w.open_files(singles, ask=False) and len(w.groups["TEST-A"]) == 5, w.groups
        magic = b"Varian UV-VIS Spectrophotometer"
        blob = bytearray([len(magic)]) + magic
        blob += bytes(0x400 - len(blob))
        for k in range(3):
            blk = bytearray(256)
            t = f"DSW-T-{0.5 / 2 ** k:g}".replace(".", "p").encode() + b"mgml"
            blk[:len(t)] = t
            blob += blk
            for x in np.arange(800.0, 249.0, -1.0):
                blob += struct.pack("<ff", x, 0.4 / 2 ** k * np.exp(-((x - 450) / 40) ** 2))
            blob += bytes(64)
        dsw = out / "input" / "test.DSW"
        dsw.write_bytes(bytes(blob))
        assert w.open_files([dsw], ask=False) and len(w.groups.get("DSW-T", [])) == 3, w.groups
        report.append("multiple files + DSW: OK")
        # Overlay (Messung + TD-DFT) und Fluoreszenz
        orca = out / "input" / "orca.out"
        orca.write_text("\n".join([
            "-" * 77, "         ABSORPTION SPECTRUM VIA TRANSITION ELECTRIC DIPOLE MOMENTS", "-" * 77,
            "State   Energy    Wavelength  fosc         T2        TX        TY        TZ",
            "        (cm-1)      (nm)                 (au**2)    (au)      (au)      (au)", "-" * 77,
            "   1   20660.0    484.0   0.150000000   1.00000   1.00000   0.00000   0.00000",
            "   2   38168.0    262.0   0.600000000   3.00000   1.70000   0.00000   0.00000", ""]),
            encoding="utf-8")
        assert len(__import__("uvvis_extra").parse_tddft(orca)["E"]) == 2
        w.ov_state["entries"] = [
            {"kind": "exp", "file": str(csv), "sample": "TEST-A-0p25mgml_THF", "label": "exp", "color": "#000000",
             "visible": True, "own_norm": False, "nlo": 450, "nhi": 520, "conc": 0.25, "unit": "mg/mL",
             "mw": 95.1, "d": 1.0},
            {"kind": "calc", "file": str(orca), "label": "calc", "color": "#1764e8", "visible": True,
             "own_norm": False, "nlo": 450, "nhi": 520, "fwhm": 0.3, "shift": 0.0}]
        w.ov_state.update(nlo=450, nhi=520)
        for mode in ("norm", "eps"):
            w.ov_state["mode"] = mode
            files = w.ov_tab.export_to(out / f"overlay_{mode}")
            assert all(f.exists() for f in files), files
        em = out / "input" / "emission.csv"
        xs = np.arange(400.0, 701.0, 1.0)
        em.write_text("\r\n".join(["EM ex352,,", "Wavelength (nm),Intensity (a.u.),"] +
                                    [f"{x:.1f},{500 * np.exp(-((x - 520) / 30) ** 2):.3f}," for x in xs]),
                      encoding="utf-8", newline="")
        w.fl_state.update(abs_file=str(csv), abs_sample="TEST-A-0p25mgml_THF", em_file=str(em),
                          em_sample="EM ex352", ex=352.0, alo=450, ahi=520, eps="1234",
                          photo_day=str(photo), bg_day="grabcut", photo_uv=str(photo), bg_uv="none")
        files = w.fl_tab.export_to(out / "fluorescence")
        assert all(f.exists() for f in files), files
        # Cary-Eclipse-CSV (Name in Zeile 2, λex aus den Metadaten)
        ecl = out / "input" / "eclipse.csv"
        ecl.write_text("\r\n".join(["Wavelength (nm),Intensity (a.u.),Z Axis,", ",PROBE-1-45p3µM", ",1"] +
                                     [f"{x:g},{300 * np.exp(-((x - 480) / 25) ** 2):.4f}" for x in xs] +
                                     ["", "PROBE-1-45p3µM", "Ex. Wavelength (nm)               387.00"]),
                       encoding="utf-8", newline="")
        es = core.read_samples([ecl])
        assert len(es) == 1 and es[0]["name"] == "PROBE-1-45p3µM" and es[0]["ex"] == 387.0, es[0]["name"]
        report.append("overlay + fluorescence + Eclipse CSV: OK")
        # Schultern, Zusatzbande, Cutoff und Ausblenden pro Bande
        w.open_csv(csv, ask=False)
        w.glob.update(shoulders=True, bands_add=True, bands="600")
        w.series["TEST-A"]["band_cutoff"] = {"484": 2.0}
        w.series["TEST-A"]["band_hidden"] = [600]
        cfg_t = w.make_cfg("TEST-A")
        _, _, res_t = core.prepare_series(cfg_t, w.groups["TEST-A"], "TEST-A")
        lams = {int(round(r["lam_requested"])): r for r in res_t}
        assert 600 in lams and lams[600].get("hidden") and lams[484]["cutoff"] == 2.0, sorted(lams)
        report.append("bands/shoulders/cutoff: OK")

        def pump(sec=0.4):
            import time
            t0 = time.time()
            while time.time() - t0 < sec:
                QApplication.processEvents()
                time.sleep(0.02)

        # TD-DFT mit Spin-Bahn-Kopplung (ORCA 6): Zustandsnamen wie "0-1.0A" dürfen nicht als Zahl gelten
        soc = out / "input" / "orca_soc.out"
        hdr = ["-" * 104, "      SOC CORRECTED ABSORPTION SPECTRUM VIA TRANSITION ELECTRIC DIPOLE MOMENTS", "-" * 104,
               "      Transition         Energy     Energy  Wavelength fosc(D2)      D2       |DX|      |DY|      |DZ|",
               "                          (eV)      (cm-1)    (nm)                 (au**2)    (au)      (au)      (au)",
               "-" * 104,
               "  0-1.0A  ->  1-3.0A    2.898266   23376.1   427.8   0.000000000   0.00000   0.00001   0.00000   0.00000",
               "  0-1.0A  -> 10-1.0A    4.004892   32301.6   309.6   0.120701528   1.23017   0.69013   0.85513   0.15048",
               ""]
        soc.write_text("\n".join(orca.read_text(encoding="utf-8").splitlines() + hdr), encoding="utf-8")
        import uvvis_extra as XX
        r_soc = XX.parse_tddft(soc)
        assert r_soc["variant"] == "soc" and abs(r_soc["E"][1] - 4.004892) < 1e-6 and abs(r_soc["f"][1] - 0.1207) < 1e-3, r_soc
        assert XX.parse_tddft(soc, "nosoc")["variant"] == "nosoc"
        report.append("TD-DFT SOC: OK")

        # Zuschnitt von Strukturen: senkrechte/waagerechte Endbindungen dürfen nicht abgeschnitten werden
        import pymupdf
        tdoc = pymupdf.open()
        pg = tdoc.new_page(width=200, height=200)
        pg.draw_line((50, 50), (150, 50), width=1)        # waagerecht
        pg.draw_line((100, 50), (100, 160), width=1)      # senkrecht, reicht am weitesten nach unten
        cr = IMG._content_rect(pg)
        assert cr.y1 >= 160 and cr.x0 <= 50 and cr.x1 >= 150, cr
        report.append("structure trimming: OK")

        # Overlay-Optionen: Skalierung, Versatz, Linienart, Strich-Modi
        w.ov_state["entries"][1].update(scale=0.8, offset=0.2, ls="dotted", lw=2.0)
        for mode in ("max", "band", "faxis"):
            w.ov_state.update(mode="norm", stick_mode=mode, sticks=True)
            files = w.ov_tab.export_to(out / f"overlay_sticks_{mode}")
            assert all(f.exists() for f in files), files
        report.append("overlay options: OK")

        # Export ins Projekt inkl. Excel-Report
        w.open_csv(csv, ask=False)
        w.series["TEST-A"]["mw"] = 95.1
        w.set_structure(str(cdxml), "TEST-A", render=False, sync=True, autofill_mw=False)
        w.render()
        pump()
        files = w.export_all()
        rep = proj / "exports" / "epsilon" / "TEST-A" / "TEST-A_report.xlsx"
        assert rep in files and rep.exists(), files
        from openpyxl import load_workbook
        wbk = load_workbook(rep)
        assert len(wbk.sheetnames) >= 3 and any("nm" in n for n in wbk.sheetnames), wbk.sheetnames
        report.append("export into project + Excel report: OK")

        # Projekt automatisch gespeichert -> in neuem Fenster öffnen (Dateien als Kopie, Pfade relativ)
        w.ov_state["entries"][0]["file"] = w.import_file(w.ov_state["entries"][0]["file"], "data")
        w.push_undo()
        pump(1.5)                                      # automatisches Speichern abwarten
        w.maybe_save()
        txt = (proj / PRJ.PROJECT_FILE).read_text(encoding="utf-8")
        assert str(out / "input") not in txt.replace(str(proj), ""), "absolute Pfade im Projekt"
        assert (proj / "data" / csv.name).exists() and (proj / "images" / cdxml.name).exists()
        w2 = MainWindow()
        assert w2.open_project(proj, ask=False)
        pump()
        assert set(w2.groups) == {"TEST-A", "TEST-B"} and w2.series["TEST-A"]["mw"] == 95.1, w2.groups
        assert len(w2.ov_state.get("entries", [])) == 2 and w2.fl_state.get("em_file"), "Tabs nicht geladen"
        report.append("project save/load: OK")

        # Undo/Redo
        before = w2.glob["cutoff"]
        w2.sp_cut.setValue(before + 0.5)
        pump(0.8)
        assert w2.glob["cutoff"] == before + 0.5
        w2.undo()
        pump()
        assert abs(w2.glob["cutoff"] - before) < 1e-9, w2.glob["cutoff"]
        w2.redo()
        pump()
        assert abs(w2.glob["cutoff"] - (before + 0.5)) < 1e-9, w2.glob["cutoff"]
        report.append("undo/redo: OK")
        for lang in ("en", "de"):
            w.switch_lang(lang)
        report.append("OK")
    except Exception:
        ok = False
        report.append("FAILED\n" + traceback.format_exc())
    (out / "selftest_report.txt").write_text("\n".join(report), encoding="utf-8")
    print("\n".join(report))
    return 0 if ok else 1


# ----------------------------------------------------------------------------
def install_crash_handler():
    def hook(t, v, tb):
        msg = "".join(traceback.format_exception(t, v, tb))
        try:
            (IMG.app_data_dir() / "crash.log").write_text(msg, encoding="utf-8")
        except Exception:
            pass
        if QApplication.instance() and os.environ.get("QT_QPA_PLATFORM") != "offscreen":
            QMessageBox.critical(None, "Error", msg[-2000:])
        sys.__excepthook__(t, v, tb)
    sys.excepthook = hook


def main():
    if "--selftest" in sys.argv:
        i = sys.argv.index("--selftest")
        outdir = sys.argv[i + 1] if len(sys.argv) > i + 1 and not sys.argv[i + 1].startswith("--") \
            else "selftest_out"
        sys.exit(selftest(outdir, "--selftest-isnet" in sys.argv))
    install_crash_handler()
    QApplication.setApplicationName(APP_NAME)
    app = QApplication(sys.argv)
    w = MainWindow()
    w.show()
    args = [a for a in sys.argv[1:] if not a.startswith("-")]
    if args and Path(args[0]).exists():
        w.open_csv(args[0])          # Datei per Doppelklick / Drag auf das Programmsymbol
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
