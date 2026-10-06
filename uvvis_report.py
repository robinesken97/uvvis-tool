"""
uvvis_report.py - Ausführlicher Excel-Report einer Verdünnungsreihe.

Blätter:
  Übersicht    Parameter, Ergebnisse aller Banden (Tool-Werte + Kontrolle aus den Bandenblättern),
               Methodik mit Formeln
  Spektren     A(λ) aller Proben (basislinienkorrigiert und roh) + Diagramm
  <λ> nm       je Bande: Messpunkte, Fit-Flag, vollständige Regression als Excel-Formeln
               (SUMMENPRODUKT), Tool-Werte daneben mit Abweichung, Residuen, Diagramm

Messwerte und Tool-Ergebnisse sind blau (feste Zahlen), alle Formeln schwarz. Ändert man in einem
Bandenblatt die Spalte „im Fit“ (1/0), rechnet Excel die Regression neu.
"""
from __future__ import annotations

import datetime as _dt
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.marker import Marker
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import uvvis_core as core
from uvvis_i18n import get_lang

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HEAD_FILL = PatternFill("solid", fgColor="E7EEF7")
KEY_FILL = PatternFill("solid", fgColor="FFF2CC")
PALETTE = ["000000", "E8231B", "1764E8", "1F9E4A", "A35CE0", "E08A00", "00A0B0", "8C564B"]


def _fit_page(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
THIN = Side(style="thin", color="B0B0B0")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

TXT = {
    "summary": ("Übersicht", "Summary"),
    "spectra": ("Spektren", "Spectra"),
    "title": ("UVVisTool – ε-Auswertung: {s}", "UVVisTool – ε analysis: {s}"),
    "params": ("Parameter", "Parameters"),
    "version": ("Programmversion", "Program version"),
    "date": ("Erstellt", "Created"),
    "files": ("Messdateien", "Measurement files"),
    "unit": ("Konzentrationseinheit", "Concentration unit"),
    "mw": ("Molmasse M [g/mol]", "Molar mass M [g/mol]"),
    "d": ("Schichtdicke d [cm]", "Path length d [cm]"),
    "fc": ("Umrechnung f_c: c [mol/L] = c × f_c", "Conversion f_c: c [mol/L] = c × f_c"),
    "cut": ("Fit-Cutoff A (global)", "Fit cutoff A (global)"),
    "minpts": ("min. Fitpunkte", "min. fit points"),
    "intercept": ("Achsenabschnitt im Fit", "Intercept in fit"),
    "yes": ("ja", "yes"), "no": ("nein (durch Ursprung)", "no (through origin)"),
    "bl": ("Basislinie", "Baseline"),
    "bl_window": ("Basislinienfenster [nm]", "Baseline window [nm]"),
    "bl_kept": ("im Fenster behaltener Anteil (a_w)", "part kept in window (a_w)"),
    "results": ("Ergebnisse", "Results"),
    "lam": ("λ [nm]", "λ [nm]"), "type": ("Typ", "Type"), "n": ("n", "n"),
    "slope": ("Steigung m", "Slope m"), "se": ("SE", "SE"), "b": ("Achsenabschnitt b", "Intercept b"),
    "r2": ("R²", "R²"), "coef_tool": ("{sym} (Tool)", "{sym} (tool)"), "se_coef": ("SE {sym}", "SE {sym}"),
    "ci": ("95 %-KI ±", "95 % CI ±"), "coef_ctrl": ("{sym} (Excel)", "{sym} (Excel)"),
    "delta": ("Δ Excel − Tool", "Δ Excel − tool"), "cutoff": ("Cutoff A", "Cutoff A"),
    "notes": ("Hinweise", "Notes"), "max": ("Maximum", "maximum"), "sh": ("Schulter", "shoulder"),
    "method": ("Methodik", "Method"),
    "sample": ("Probe", "Sample"), "conc": ("c [{u}]", "c [{u}]"), "A": ("A", "A"),
    "Ad": ("A/d [cm⁻¹]", "A/d [cm⁻¹]"), "w": ("im Fit (1/0)", "in fit (1/0)"),
    "w_tool": ("Tool: im Fit", "Tool: in fit"), "res": ("Residuum", "Residual"),
    "fit": ("Regression (Excel-Formeln)", "Regression (Excel formulas)"),
    "tool": ("Tool", "Tool"), "band": ("Bande {l} nm ({t})", "Band {l} nm ({t})"),
    "chart_fit": ("A/d gegen c – {l} nm", "A/d vs c – {l} nm"),
    "chart_spec": ("Spektren (basislinienkorrigiert)", "Spectra (baseline-corrected)"),
    "wl": ("Wellenlänge [nm]", "Wavelength [nm]"),
    "corr": ("A (korrigiert)", "A (corrected)"), "raw": ("A (roh)", "A (raw)"),
    "used": ("im Fit", "in fit"), "excl": ("ausgeschlossen", "excluded"), "line": ("Regression", "Regression"),
    "edit_hint": ("Spalte „im Fit“ (1/0) kann geändert werden – die Regression rechnet neu.",
                  "Column “in fit” (1/0) can be edited – the regression recalculates."),
    "blue_hint": ("Blau = Messwerte bzw. Tool-Ergebnisse (feste Zahlen), schwarz = Excel-Formeln.",
                  "Blue = measured values or tool results (fixed numbers), black = Excel formulas."),
}


def t(key, **kw):
    de, en = TXT[key]
    s = en if get_lang() == "en" else de
    return s.format(**kw) if kw else s


METHOD = {
    "de": [
        "Lambert-Beer:  A = ε · c · d   ⇒   A/d = ε · c (+ b)",
        "Für jede Bande wird A/d gegen c linear angepasst (kleinste Quadrate):",
        "  x̄ = Σw·x / n,  ȳ = Σw·y / n,  Sxx = Σw·(x−x̄)²,  Sxy = Σw·(x−x̄)(y−ȳ)",
        "  m = Sxy / Sxx,   b = ȳ − m·x̄   (ohne Achsenabschnitt: m = Σw·x·y / Σw·x², b = 0)",
        "  RSS = Σw·(y − m·x − b)²,   s = √(RSS / (n − p)),   p = 2 (bzw. 1 ohne Achsenabschnitt)",
        "  SE(m) = s / √Sxx,   SE(b) = s · √(1/n + x̄²/Sxx),   R² = 1 − RSS / Σw·(y−ȳ)²",
        "ε = m / f_c   mit f_c = 10⁻³ (mM), 10⁻⁶ (µM), 1 (M) bzw. 1/M (mg/mL; ohne M: spezifischer Koeffizient a)",
        "SE(ε) = SE(m) / f_c;   95 %-Konfidenzintervall: ± t(0,975; n − p) · SE(ε)",
        "w = 1, wenn der Punkt in den Fit eingeht (A ≤ Cutoff der Bande), sonst 0.",
        "Basislinie (automatisch, Reihe): Fenster dort, wo die konzentrierteste Probe am wenigsten absorbiert;",
        "  abgezogen wird je Probe Offset_i = A_i(Fenster) − a_w·c_i (der zu c proportionale Anteil a_w bleibt).",
    ],
    "en": [
        "Beer–Lambert:  A = ε · c · d   ⇒   A/d = ε · c (+ b)",
        "For every band, A/d is fitted linearly against c (least squares):",
        "  x̄ = Σw·x / n,  ȳ = Σw·y / n,  Sxx = Σw·(x−x̄)²,  Sxy = Σw·(x−x̄)(y−ȳ)",
        "  m = Sxy / Sxx,   b = ȳ − m·x̄   (without intercept: m = Σw·x·y / Σw·x², b = 0)",
        "  RSS = Σw·(y − m·x − b)²,   s = √(RSS / (n − p)),   p = 2 (or 1 without intercept)",
        "  SE(m) = s / √Sxx,   SE(b) = s · √(1/n + x̄²/Sxx),   R² = 1 − RSS / Σw·(y−ȳ)²",
        "ε = m / f_c   with f_c = 10⁻³ (mM), 10⁻⁶ (µM), 1 (M) or 1/M (mg/mL; without M: specific coefficient a)",
        "SE(ε) = SE(m) / f_c;   95 % confidence interval: ± t(0.975; n − p) · SE(ε)",
        "w = 1 if the point enters the fit (A ≤ cutoff of the band), else 0.",
        "Baseline (automatic, series): window where the most concentrated sample absorbs least;",
        "  per sample Offset_i = A_i(window) − a_w·c_i is subtracted (the part a_w proportional to c is kept).",
    ],
}


def _set(ws, ref, value, font=BLACK, fmt=None, fill=None, bold=False):
    c = ws[ref]
    c.value = value
    c.font = BOLD if bold else font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    return c


def _header(ws, row, col, labels):
    for j, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=lab)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _num(v):
    try:
        v = float(v)
        return v if np.isfinite(v) else None
    except (TypeError, ValueError):
        return None


def _sheet_name(lam, used):
    name = f"{lam:g} nm"
    k = 2
    while name in used:
        name = f"{lam:g} nm ({k})"
        k += 1
    used.add(name)
    return name


def write_report(path, cfg, raw_spectra, spectra, results, series, files=(), version=""):
    """raw_spectra: Spektren vor der Basislinie; spectra/results/cfg aus core.prepare_series."""
    sym, _unit = core.coeff_symbol_unit(cfg)
    sym = "ε" if sym == "ε" else "a"
    coef_unit = "M⁻¹ cm⁻¹" if sym == "ε" else "L g⁻¹ cm⁻¹"
    fac = core.conc_factor(cfg) or 1.0
    unit = cfg["conc_unit"]
    unit_disp = "µM" if unit == "uM" else unit
    wb = Workbook()
    ws = wb.active
    ws.title = t("summary")

    # ---------------- Übersicht: Parameter
    _set(ws, "A1", t("title", s=series), TITLE)
    _set(ws, "A3", t("params"), bold=True)
    bl = cfg.get("_baseline_info") or {}
    rows = [
        (t("version"), version, None),
        (t("date"), _dt.datetime.now().strftime("%Y-%m-%d %H:%M"), None),
        (t("files"), ", ".join(Path(f).name for f in files) or "–", None),
        (t("unit"), unit_disp, None),
        (t("mw"), _num(cfg.get("molar_mass_g_mol")), "0.00"),
        (t("d"), float(cfg["path_length_cm"]), "0.000"),
        (t("fc"), None, "0.000E+00"),
        (t("cut"), float(cfg["max_abs_fit"]), "0.00"),
        (t("minpts"), int(cfg["min_fit_points"]), "0"),
        (t("intercept"), t("yes") if cfg["fit_intercept"] else t("no"), None),
        (t("bl"), f"{cfg.get('baseline_nm')} / {cfg.get('baseline_mode')}", None),
        (t("bl_window"), (f"{bl['window'][0]:.0f} – {bl['window'][1]:.0f}" if bl else "–"), None),
        (t("bl_kept"), _num(bl.get("kept_slope")) if bl else None, "0.0000"),
    ]
    prow = {}
    for i, (lab, val, fmt) in enumerate(rows):
        r = 4 + i
        _set(ws, f"A{r}", lab)
        _set(ws, f"B{r}", val, BLUE, fmt)
        prow[lab] = r
    r_mw, r_d, r_fc = prow[t("mw")], prow[t("d")], prow[t("fc")]
    # f_c als Formel (mg/mL über die Molmasse)
    if unit == "mg/mL":
        ws[f"B{r_fc}"] = f'=IF(ISNUMBER(B{r_mw}),1/B{r_mw},1)'
    else:
        ws[f"B{r_fc}"] = {"mM": 1e-3, "uM": 1e-6, "M": 1.0}.get(unit, 1.0)
        ws[f"B{r_fc}"].font = BLUE
    ws[f"B{r_fc}"].number_format = "0.000E+00"
    ws[f"B{r_fc}"].font = BLACK if unit == "mg/mL" else BLUE
    ws[f"C{r_fc}"] = ("1/M (Formel)" if unit == "mg/mL" else "")
    if bl and bl.get("offsets") is not None:
        r0 = 4 + len(rows) + 1
        _set(ws, f"A{r0}", "Offsets" if get_lang() == "en" else "Offsets je Probe", bold=True)
        for k, (s_, o) in enumerate(sorted(zip(spectra, bl["offsets"]), key=lambda p: -p[0]["c"])):
            _set(ws, f"A{r0 + 1 + k}", s_.get("name", f"#{k + 1}"))
            _set(ws, f"B{r0 + 1 + k}", float(o), BLUE, "0.00000")
        res_row = r0 + len(spectra) + 3
    else:
        res_row = 4 + len(rows) + 2

    # ---------------- Bandenblätter
    used_names = {ws.title, t("spectra")}
    band_links = []
    for r in results:
        name = _sheet_name(r["lam"], used_names)
        bs = wb.create_sheet(name)
        link = _band_sheet(bs, r, cfg, spectra, unit_disp, sym, coef_unit, ws.title, r_d, r_fc)
        band_links.append((name, link))

    # ---------------- Übersicht: Ergebnisse
    _set(ws, f"A{res_row}", t("results"), bold=True)
    hdr = [t("lam"), t("type"), t("n"), t("slope"), t("se"), t("b"), t("se"), t("r2"),
           t("coef_tool", sym=sym) + f" [{coef_unit}]", t("se_coef", sym=sym), t("ci"),
           t("coef_ctrl", sym=sym), t("delta"), t("cutoff"), t("notes")]
    _header(ws, res_row + 1, 1, hdr)
    for i, (r, (name, link)) in enumerate(zip(results, band_links)):
        rr = res_row + 2 + i
        ft = r["fit"]
        vals = [r["lam"], t("sh") if r.get("shoulder") else t("max"), ft["n"] if ft else len(r["used"]),
                ft and ft["slope"], ft and ft["se_slope"], ft and ft["intercept"], ft and ft["se_intercept"],
                ft and ft["r2"], r.get("eps"), r.get("eps_err"), r.get("eps_ci95")]
        fmts = ["0", None, "0", "0.0000", "0.0000", "0.0000", "0.0000", "0.00000", "0", "0.0", "0.0"]
        if sym == "a":
            fmts[8:11] = ["0.0000", "0.0000", "0.0000"]
        for j, (v, f) in enumerate(zip(vals, fmts)):
            c = ws.cell(row=rr, column=1 + j, value=_num(v) if j != 1 else v)
            c.font = BLUE
            c.border = BOX
            if f:
                c.number_format = f
        q = f"'{name}'!{link}"
        c = ws.cell(row=rr, column=12, value=f"=IFERROR({q},\"–\")")
        c.number_format = fmts[8]
        c.border = BOX
        c = ws.cell(row=rr, column=13, value=f"=IF(AND(ISNUMBER({q}),ISNUMBER(I{rr})),{q}-I{rr},\"–\")")
        c.number_format = "0.000E+00"
        c.border = BOX
        c = ws.cell(row=rr, column=14, value=_num(r.get("cutoff")))
        c.font = BLUE
        c.number_format = "0.00"
        c.border = BOX
        ws.cell(row=rr, column=15, value="; ".join(r.get("warnings") or [])).font = BLACK
    mrow = res_row + len(results) + 3
    _set(ws, f"A{mrow}", t("method"), bold=True)
    for k, line in enumerate(METHOD["en" if get_lang() == "en" else "de"]):
        _set(ws, f"A{mrow + 1 + k}", line)
    _set(ws, f"A{mrow + len(METHOD['de']) + 2}", t("blue_hint"))
    ws.column_dimensions["A"].width = 38
    for col in range(2, 16):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.column_dimensions["O"].width = 60
    ws.freeze_panes = "A2"
    _fit_page(ws)

    # ---------------- Spektren
    _spectra_sheet(wb.create_sheet(t("spectra"), 1), raw_spectra, spectra)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _band_sheet(bs, r, cfg, spectra, unit, sym, coef_unit, summary, r_d, r_fc):
    """Bandenblatt. Gibt die Zelle mit ε (Excel) zurück."""
    lam = r["lam"]
    _set(bs, "A1", t("band", l=f"{lam:g}", t=t("sh") if r.get("shoulder") else t("max")), TITLE)
    _set(bs, "A2", t("edit_hint"))
    # Datentabelle
    hdr = [t("sample"), t("conc", u=unit), t("A"), t("Ad"), t("w"), t("w_tool"), t("res")]
    _header(bs, 4, 1, hdr)
    used = {round(c, 12) for c, _ in r["used"]}
    pts = []
    for s_ in sorted(spectra, key=lambda s_: s_["c"]):
        A = core.abs_at(s_["x"], s_["y"], lam, cfg["avg_window_nm"])
        pts.append((s_.get("name", ""), s_["c"], A, round(s_["c"], 12) in used and np.isfinite(A)))
    first = 5
    last = first + len(pts) - 1
    cut = float(r.get("cutoff", cfg["max_abs_fit"]))
    # Parameterzellen rechts
    _set(bs, "J4", t("params"), bold=True)
    _set(bs, "J5", "d [cm]")
    bs["K5"] = f"='{summary}'!B{r_d}"
    bs["K5"].number_format = "0.000"
    _set(bs, "J6", "f_c")
    bs["K6"] = f"='{summary}'!B{r_fc}"
    bs["K6"].number_format = "0.000E+00"
    _set(bs, "J7", t("cutoff"))
    _set(bs, "K7", cut, BLUE, "0.00")
    _set(bs, "J8", t("minpts"))
    _set(bs, "K8", int(cfg["min_fit_points"]), BLUE, "0")
    for i, (name, c, A, u) in enumerate(pts):
        row = first + i
        bs.cell(row=row, column=1, value=name).font = BLACK
        _set(bs, f"B{row}", float(c), BLUE, "0.000000")
        _set(bs, f"C{row}", _num(A), BLUE, "0.00000")
        bs[f"D{row}"] = f"=C{row}/$K$5"
        bs[f"D{row}"].number_format = "0.00000"
        bs[f"E{row}"] = f'=IF(AND(ISNUMBER(C{row}),C{row}<=$K$7),1,0)'
        bs[f"E{row}"].fill = KEY_FILL
        bs[f"E{row}"].comment = Comment(t("edit_hint"), "UVVisTool")
        _set(bs, f"F{row}", 1 if u else 0, BLUE, "0")
        bs[f"G{row}"] = f'=IF(E{row}=1,D{row}-($K$12*B{row}+$K$13),"")'
        bs[f"G{row}"].number_format = "0.00000"
        for col in "ABCDEFG":
            bs[f"{col}{row}"].border = BOX
    X, Y, W = f"$B${first}:$B${last}", f"$D${first}:$D${last}", f"$E${first}:$E${last}"
    icpt = bool(cfg["fit_intercept"])
    p = 2 if icpt else 1
    # Regression als Formeln
    _set(bs, "J9", t("fit"), bold=True)
    _set(bs, "L9", t("tool"), bold=True)
    _set(bs, "M9", t("delta"), bold=True)
    ft = r["fit"] or {}
    lines = [
        ("n", f"=SUM({W})", None, "0"),
        ("x̄", f"=IF(K10>0,SUMPRODUCT({W},{X})/K10,\"\")", None, "0.000000"),
        ("m", None, ft.get("slope"), "0.000000"),
        ("b", None, ft.get("intercept"), "0.000000"),
        ("ȳ", f"=IF(K10>0,SUMPRODUCT({W},{Y})/K10,\"\")", None, "0.000000"),
        ("Sxx", f"=SUMPRODUCT({W},({X}-K11)^2)", None, "0.000000"),
        ("Sxy", f"=SUMPRODUCT({W},({X}-K11),({Y}-K14))", None, "0.000000"),
        ("RSS", f"=SUMPRODUCT({W},({Y}-(K12*{X}+K13))^2)", ft.get("rss"), "0.000E+00"),
        ("s", f"=IF(K10>{p},SQRT(K17/(K10-{p})),\"\")", None, "0.000E+00"),
        ("SE(m)", None, ft.get("se_slope"), "0.000000"),
        ("SE(b)", None, ft.get("se_intercept") if icpt else None, "0.000000"),
        ("R²", f"=IF(K10>1,1-K17/SUMPRODUCT({W},({Y}-K14)^2),\"\")", ft.get("r2"), "0.000000"),
        (f"{sym} [{coef_unit}]", "=IF(K10>=K8,IFERROR(K12/K6,\"\"),\"\")", r.get("eps"),
         "0" if sym == "ε" else "0.0000"),
        (f"SE({sym})", "=IFERROR(K19/K6,\"\")", r.get("eps_err"), "0.0" if sym == "ε" else "0.0000"),
        ("t(0.975; n−p)", f"=IF(K10>{p},TINV(0.05,K10-{p}),\"\")", r.get("t95"), "0.0000"),
        ("95 % CI ±", "=IFERROR(K24*K23,\"\")", r.get("eps_ci95"), "0.0" if sym == "ε" else "0.0000"),
    ]
    # Steigung/Achsenabschnitt/SE abhängig vom Modell (Zeilen 12, 13, 19, 20)
    if icpt:
        f_m, f_b = "=IFERROR(K16/K15,\"\")", "=IFERROR(K14-K12*K11,\"\")"
        f_sem = "=IFERROR(K18/SQRT(K15),\"\")"
        f_seb = "=IFERROR(K18*SQRT(1/K10+K11^2/K15),\"\")"
    else:
        f_m = f"=IFERROR(SUMPRODUCT({W},{X},{Y})/SUMPRODUCT({W},{X},{X}),\"\")"
        f_b = "=0"
        f_sem = f"=IFERROR(K18/SQRT(SUMPRODUCT({W},{X},{X})),\"\")"
        f_seb = "=\"–\""
    for k, (lab, f, tool, fmt) in enumerate(lines):
        row = 10 + k
        if lab == "m":
            f = f_m
        elif lab == "b":
            f = f_b
        elif lab == "SE(m)":
            f = f_sem
        elif lab == "SE(b)":
            f = f_seb
        _set(bs, f"J{row}", lab)
        c = bs[f"K{row}"]
        c.value = f
        c.number_format = fmt
        c.border = BOX
        if tool is not None and _num(tool) is not None:
            _set(bs, f"L{row}", _num(tool), BLUE, fmt)
            bs[f"M{row}"] = f'=IFERROR(K{row}-L{row},"")'
            bs[f"M{row}"].number_format = "0.000E+00"
        if lab.startswith(sym + " ["):
            c.fill = KEY_FILL
    # Diagrammdaten (Werte vom Tool) + Regressionsgerade aus den Formelzellen
    r0 = max(last + 4, 29)                            # unterhalb des Parameter-/Regressionsblocks
    _set(bs, "A" + str(r0 - 1), t("used") + " / " + t("excl") + " / " + t("line"), bold=True)
    _header(bs, r0, 1, ["c (" + t("used") + ")", "A/d", "c (" + t("excl") + ")", "A/d", "c (" + t("line") + ")", "A/d"])
    d = float(cfg["path_length_cm"])
    u_pts = [(c, A / d) for _, c, A, u in pts if u]
    e_pts = [(c, A / d) for _, c, A, u in pts if not u and np.isfinite(A)]
    xs = [c for _, c, A, _u in pts]
    for i, (c, y) in enumerate(u_pts):
        _set(bs, f"A{r0 + 1 + i}", float(c), BLUE, "0.000000")
        _set(bs, f"B{r0 + 1 + i}", float(y), BLUE, "0.00000")
    for i, (c, y) in enumerate(e_pts):
        _set(bs, f"C{r0 + 1 + i}", float(c), BLUE, "0.000000")
        _set(bs, f"D{r0 + 1 + i}", float(y), BLUE, "0.00000")
    for i, xv in enumerate([0.0, max(xs) if xs else 1.0]):
        _set(bs, f"E{r0 + 1 + i}", float(xv), BLUE, "0.000000")
        bs[f"F{r0 + 1 + i}"] = f'=IFERROR($K$12*E{r0 + 1 + i}+$K$13,0)'
        bs[f"F{r0 + 1 + i}"].number_format = "0.00000"
    ch = ScatterChart()
    ch.title = t("chart_fit", l=f"{lam:g}")
    ch.style = 13
    ch.x_axis.title = f"c [{unit}]"
    ch.y_axis.title = "A/d [cm⁻¹]"
    ch.x_axis.number_format = "0.00"
    ch.y_axis.number_format = "0.00"
    ch.x_axis.scaling.min = 0
    ch.y_axis.scaling.min = 0
    ch.height, ch.width = 8, 14
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    if u_pts:
        s1 = Series(Reference(bs, min_col=2, min_row=r0 + 1, max_row=r0 + len(u_pts)),
                    Reference(bs, min_col=1, min_row=r0 + 1, max_row=r0 + len(u_pts)), title=t("used"))
        s1.marker = Marker(symbol="square", size=7)
        s1.marker.graphicalProperties.solidFill = "000000"
        s1.marker.graphicalProperties.line.solidFill = "000000"
        s1.graphicalProperties.line.noFill = True
        ch.series.append(s1)
    if e_pts:
        s2 = Series(Reference(bs, min_col=4, min_row=r0 + 1, max_row=r0 + len(e_pts)),
                    Reference(bs, min_col=3, min_row=r0 + 1, max_row=r0 + len(e_pts)), title=t("excl"))
        s2.marker = Marker(symbol="square", size=7)
        s2.marker.graphicalProperties.noFill = True
        s2.marker.graphicalProperties.line.solidFill = "7F7F7F"
        s2.graphicalProperties.line.noFill = True
        ch.series.append(s2)
    s3 = Series(Reference(bs, min_col=6, min_row=r0 + 1, max_row=r0 + 2),
                Reference(bs, min_col=5, min_row=r0 + 1, max_row=r0 + 2), title=t("line"))
    s3.marker = Marker(symbol="none")
    s3.graphicalProperties.line.solidFill = "E8231B"
    s3.graphicalProperties.line.width = 19050
    ch.series.append(s3)
    bs.add_chart(ch, "O4")
    _fit_page(bs)
    bs.column_dimensions["A"].width = 30
    for col in "BCDEFG":
        bs.column_dimensions[col].width = 13
    bs.column_dimensions["J"].width = 20
    for col in "KLM":
        bs.column_dimensions[col].width = 15
    return "K22"                                      # ε (Excel)


def _spectra_sheet(ss, raw_spectra, spectra):
    order = sorted(range(len(spectra)), key=lambda i: -spectra[i]["c"])
    xs = spectra[order[0]]["x"] if spectra else np.array([])
    n = len(spectra)
    _header(ss, 1, 1, [t("wl")] + [f"{t('corr')}: {spectra[i].get('name', '')}" for i in order]
            + [f"{t('raw')}: {spectra[i].get('name', '')}" for i in order])
    raw_by_name = {s_.get("name"): s_ for s_ in raw_spectra}
    for k, x in enumerate(xs):
        row = 2 + k
        ss.cell(row=row, column=1, value=float(x)).font = BLUE
        for j, i in enumerate(order):
            s_ = spectra[i]
            v = float(np.interp(x, s_["x"], s_["y"]))
            c = ss.cell(row=row, column=2 + j, value=round(v, 6))
            c.font = BLUE
            rs = raw_by_name.get(s_.get("name"))
            if rs is not None:
                c = ss.cell(row=row, column=2 + n + j, value=round(float(np.interp(x, rs["x"], rs["y"])), 6))
                c.font = BLUE
    ch = ScatterChart()
    ch.title = t("chart_spec")
    ch.style = 13
    ch.x_axis.title = t("wl")
    ch.y_axis.title = "A"
    ch.height, ch.width = 10, 18
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 1.2
    ch.x_axis.number_format = "0"
    ch.y_axis.number_format = "0.0"
    last = 1 + len(xs)
    for j in range(n):
        s_ = Series(Reference(ss, min_col=2 + j, min_row=1, max_row=last),
                    Reference(ss, min_col=1, min_row=2, max_row=last), title_from_data=True)
        s_.marker = Marker(symbol="none")
        s_.smooth = False
        s_.graphicalProperties.line.solidFill = PALETTE[j % len(PALETTE)]
        s_.graphicalProperties.line.width = 12700
        ch.series.append(s_)
    ss.add_chart(ch, get_column_letter(2 * n + 3) + "2")
    ss.column_dimensions["A"].width = 14
    ss.freeze_panes = "B2"
